"""
Unit tests for the excel_formatting module - annotations analysis sheet formatting.

This test suite validates the Annotations Analysis sheet formatting including:
- Auto-filter application
- Frozen panes
- Column widths
- Text wrapping
- RGB color fills
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class TestAnnotationsAnalysisFormatting:
    """Test suite for _format_annotations_analysis_sheet function."""

    def test_annotations_analysis_auto_filter(self) -> None:
        """Test that auto-filter is applied to Annotations Analysis sheet."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers and data
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row
        ws.cell(row=2, column=1, value="Sample Text")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=255)
        ws.cell(row=2, column=5, value=0)
        ws.cell(row=2, column=6, value=0)
        ws.cell(row=2, column=7, value="")
        ws.cell(row=2, column=8, value=3)

        _format_annotations_analysis_sheet(wb)

        # Verify auto-filter
        assert ws.auto_filter is not None
        assert ws.auto_filter.ref == "A1:H2"

    def test_annotations_analysis_frozen_panes(self) -> None:
        """Test that frozen panes are applied at A2."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        _format_annotations_analysis_sheet(wb)

        # Verify frozen panes
        assert ws.freeze_panes is not None
        assert ws.freeze_panes == "A2"

    def test_annotations_analysis_column_widths(self) -> None:
        """Test that column widths are correctly set."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        _format_annotations_analysis_sheet(wb)

        # Verify column widths
        assert ws.column_dimensions["A"].width == 60  # annotation_contents
        assert ws.column_dimensions["B"].width == 15  # annotation_type
        assert ws.column_dimensions["C"].width == 25  # annotation_layer_name
        assert ws.column_dimensions["D"].width == 12  # annotation_color_r
        assert ws.column_dimensions["E"].width == 12  # annotation_color_g
        assert ws.column_dimensions["F"].width == 12  # annotation_color_b
        assert ws.column_dimensions["G"].width == 12  # annotation_color_sample
        assert ws.column_dimensions["H"].width == 20  # annotation_count

    def test_annotations_analysis_text_wrapping(self) -> None:
        """Test that text wrapping is enabled on annotation_contents column."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers and data
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data rows
        ws.cell(row=2, column=1, value="Long text content that should wrap")
        ws.cell(row=3, column=1, value="Another long text")

        _format_annotations_analysis_sheet(wb)

        # Verify text wrapping on column A (annotation_contents)
        assert ws.cell(row=2, column=1).alignment.wrap_text is True
        assert ws.cell(row=3, column=1).alignment.wrap_text is True

    def test_annotations_analysis_color_fill(self) -> None:
        """Test that RGB color fills are applied to annotation_color_sample column."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with RGB values
        ws.cell(row=2, column=1, value="Red Text")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=255)  # R
        ws.cell(row=2, column=5, value=0)  # G
        ws.cell(row=2, column=6, value=0)  # B
        ws.cell(row=2, column=7, value="")  # color_sample
        ws.cell(row=2, column=8, value=1)

        # Add another row with different color
        ws.cell(row=3, column=1, value="Green Text")
        ws.cell(row=3, column=2, value="MTEXT")
        ws.cell(row=3, column=3, value="Layer2")
        ws.cell(row=3, column=4, value=0)  # R
        ws.cell(row=3, column=5, value=255)  # G
        ws.cell(row=3, column=6, value=0)  # B
        ws.cell(row=3, column=7, value="")  # color_sample
        ws.cell(row=3, column=8, value=2)

        _format_annotations_analysis_sheet(wb)

        # Verify color fills on column G (annotation_color_sample)
        # Row 2 should have red fill (FF0000)
        cell_fill_2 = ws.cell(row=2, column=7).fill
        assert cell_fill_2.start_color.rgb == "00FF0000"  # openpyxl uses ARGB format

        # Row 3 should have green fill (00FF00)
        cell_fill_3 = ws.cell(row=3, column=7).fill
        assert cell_fill_3.start_color.rgb == "0000FF00"

    def test_annotations_analysis_empty_sheet(self) -> None:
        """Test formatting with headers only (no data rows)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add only headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Should not raise exception
        _format_annotations_analysis_sheet(wb)

        # Verify basic formatting still applied
        assert ws.freeze_panes == "A2"
        assert ws.column_dimensions["A"].width == 60

    def test_annotations_analysis_invalid_rgb_values(self) -> None:
        """Test handling of invalid RGB values (out of range)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with invalid RGB values (should be clamped or skipped)
        ws.cell(row=2, column=1, value="Invalid Color")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=300)  # Invalid R (> 255)
        ws.cell(row=2, column=5, value=-10)  # Invalid G (< 0)
        ws.cell(row=2, column=6, value=100)  # Valid B
        ws.cell(row=2, column=7, value="")
        ws.cell(row=2, column=8, value=1)

        # Should not raise exception
        _format_annotations_analysis_sheet(wb)

    def test_annotations_analysis_rgb_boundary_values(self) -> None:
        """Test RGB color fills with boundary values (0 and 255)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add row with white color (255, 255, 255)
        ws.cell(row=2, column=1, value="White Text")
        ws.cell(row=2, column=4, value=255)
        ws.cell(row=2, column=5, value=255)
        ws.cell(row=2, column=6, value=255)
        ws.cell(row=2, column=7, value="")

        # Add row with black color (0, 0, 0)
        ws.cell(row=3, column=1, value="Black Text")
        ws.cell(row=3, column=4, value=0)
        ws.cell(row=3, column=5, value=0)
        ws.cell(row=3, column=6, value=0)
        ws.cell(row=3, column=7, value="")

        _format_annotations_analysis_sheet(wb)

        # Verify white fill
        white_fill = ws.cell(row=2, column=7).fill
        assert white_fill.start_color.rgb == "00FFFFFF"

        # Verify black fill
        black_fill = ws.cell(row=3, column=7).fill
        assert black_fill.start_color.rgb == "00000000"
