"""
Unit tests for the Instructions sheet formatting in excel_formatting module.

This test suite validates:
- Column widths are set correctly
- Bold font is applied to section headers
- Text wrapping is enabled on description column
- Color fills are applied to color sample cells
- NO auto-filter is applied
- NO freeze panes are applied
- Empty and missing sheet handling
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import (
    EXCEL_FILL_COLOR_NESTED_BLOCK,
    EXCEL_FILL_COLOR_SCALE_NEGATIVE,
    EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE,
    EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
    EXCEL_SHEET_INSTRUCTIONS,
)
from core.excel_formatting import _format_instructions_sheet


class TestInstructionsSheetFormatting:
    """Test suite for _format_instructions_sheet function."""

    def _create_instructions_workbook(self) -> Workbook:
        """Create a workbook with Instructions sheet populated with expected content."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_INSTRUCTIONS

        # Add header row
        ws.append(["Identifier", "Description"])

        # Color Coding section
        ws.append(["Color Coding", ""])
        ws.append(["Yellow", "Scale variance detected"])
        ws.append(["Orange", "Negative scale detected"])
        ws.append(["Red", "Scale variance with negatives"])
        ws.append(["Light Green", "Nested block"])

        # Separator
        ws.append(["", ""])

        # Sheet Descriptions section
        ws.append(["Sheet Descriptions", ""])
        ws.append(["All Blocks", "Consolidated block-centric view"])
        ws.append(["Block Analysis", "Simplified inventory"])

        return wb

    def test_format_instructions_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]
        assert ws.column_dimensions["A"].width == 20
        assert ws.column_dimensions["B"].width == 60

    def test_format_instructions_bold_section_headers(self, temp_dir: str) -> None:
        """Test that section headers have bold font."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 2: "Color Coding" should be bold
        color_coding_cell = ws.cell(row=2, column=1)
        assert color_coding_cell.font.bold is True

        # Row 8: "Sheet Descriptions" should be bold
        sheet_desc_cell = ws.cell(row=8, column=1)
        assert sheet_desc_cell.font.bold is True

    def test_format_instructions_text_wrapping(self, temp_dir: str) -> None:
        """Test that description column has text wrapping."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Check text wrapping on description cells
        for row_idx in range(2, ws.max_row + 1):
            desc_cell = ws.cell(row=row_idx, column=2)
            assert desc_cell.alignment is not None
            assert desc_cell.alignment.wrap_text is True

    def test_format_instructions_no_autofilter(self, temp_dir: str) -> None:
        """Test that NO auto-filter is applied to Instructions sheet."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # auto_filter.ref should be None or empty
        assert ws.auto_filter.ref is None or ws.auto_filter.ref == ""

    def test_format_instructions_no_freeze_panes(self, temp_dir: str) -> None:
        """Test that NO freeze panes are applied to Instructions sheet."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # freeze_panes should be None
        assert ws.freeze_panes is None

    def test_format_instructions_yellow_fill(self, temp_dir: str) -> None:
        """Test that Yellow row has correct color fill."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 3 (Yellow) should have yellow fill in column B
        yellow_cell = ws.cell(row=3, column=2)
        assert yellow_cell.fill is not None
        assert yellow_cell.fill.fill_type == "solid"
        assert yellow_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

    def test_format_instructions_orange_fill(self, temp_dir: str) -> None:
        """Test that Orange row has correct color fill."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 4 (Orange) should have orange fill in column B
        orange_cell = ws.cell(row=4, column=2)
        assert orange_cell.fill is not None
        assert orange_cell.fill.fill_type == "solid"
        assert orange_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE

    def test_format_instructions_red_fill(self, temp_dir: str) -> None:
        """Test that Red row has correct color fill."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 5 (Red) should have red fill in column B
        red_cell = ws.cell(row=5, column=2)
        assert red_cell.fill is not None
        assert red_cell.fill.fill_type == "solid"
        assert red_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

    def test_format_instructions_green_fill(self, temp_dir: str) -> None:
        """Test that Light Green row has correct color fill."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 6 (Light Green) should have green fill in column B
        green_cell = ws.cell(row=6, column=2)
        assert green_cell.fill is not None
        assert green_cell.fill.fill_type == "solid"
        assert green_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_NESTED_BLOCK

    def test_format_instructions_all_color_fills(self, temp_dir: str) -> None:
        """Test that all four color fills are applied correctly."""
        wb = self._create_instructions_workbook()
        _format_instructions_sheet(wb)

        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        expected_fills = [
            (3, EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE),   # Yellow
            (4, EXCEL_FILL_COLOR_SCALE_NEGATIVE),            # Orange
            (5, EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE),   # Red
            (6, EXCEL_FILL_COLOR_NESTED_BLOCK),              # Light Green
        ]

        for row_idx, expected_color in expected_fills:
            cell = ws.cell(row=row_idx, column=2)
            assert cell.fill is not None
            assert cell.fill.fill_type == "solid"
            assert cell.fill.start_color.rgb == expected_color

    def test_format_instructions_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Instructions sheet is handled gracefully."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_INSTRUCTIONS

        # Apply formatting (should not crash on empty sheet)
        _format_instructions_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 20
        assert ws.column_dimensions["B"].width == 60

    def test_format_instructions_missing_sheet(self, temp_dir: str) -> None:
        """Test that missing Instructions sheet is handled gracefully."""
        wb = Workbook()
        # Default sheet has different name, so Instructions doesn't exist

        # Should not raise an exception
        _format_instructions_sheet(wb)

        # Verify default sheet is unchanged
        assert len(wb.sheetnames) == 1
        assert EXCEL_SHEET_INSTRUCTIONS not in wb.sheetnames
