"""
Unit tests for the excel_formatting module - Block Definitions sheet.

This test suite validates Block Definitions sheet formatting including:
- Auto-filter application
- Column width settings
- Frozen panes
- Green highlighting for nested blocks
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import (
    EXCEL_FILL_COLOR_NESTED_BLOCK,
    EXCEL_SHEET_BLOCK_DEFINITIONS,
)
from core.excel_formatting import _format_block_definitions_sheet


class TestBlockDefinitionsFormatting:
    """Test suite for _format_block_definitions_sheet function."""

    def test_format_block_definitions_autofilter(self) -> None:
        """Test that auto-filter is applied to Block Definitions sheet."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )
        ws.append(["VALVE", "VALVE", "Inserted", False, "", 8])

        _format_block_definitions_sheet(wb)

        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:F2"

    def test_format_block_definitions_column_widths(self) -> None:
        """Test that column widths are set correctly."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )

        _format_block_definitions_sheet(wb)

        assert ws.column_dimensions["A"].width == 30  # block_raw_name
        assert ws.column_dimensions["B"].width == 30  # block_resolved_name
        assert ws.column_dimensions["C"].width == 20  # block_insertion_status
        assert ws.column_dimensions["D"].width == 15  # block_is_nested
        assert ws.column_dimensions["E"].width == 40  # block_nested_parent_names
        assert ws.column_dimensions["F"].width == 20  # block_entity_count

    def test_format_block_definitions_frozen_panes(self) -> None:
        """Test that frozen panes are applied at B2 (header row and first column)."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )

        _format_block_definitions_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_format_block_definitions_green_highlighting_nested(self) -> None:
        """Test that nested blocks are highlighted in green."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )
        ws.append(["VALVE", "VALVE", "Inserted", False, "", 8])
        ws.append(["PIPE", "PIPE", "Inserted", True, "VALVE", 12])
        ws.append(["TAG", "TAG", "Nested Only", True, "VALVE, PIPE", 4])

        _format_block_definitions_sheet(wb)

        # Row 2 (VALVE) should NOT have green fill (not nested)
        valve_fill = ws.cell(row=2, column=1).fill
        if valve_fill.fill_type == "solid":
            assert valve_fill.start_color.rgb != EXCEL_FILL_COLOR_NESTED_BLOCK

        # Row 3 (PIPE) should have green fill (nested)
        assert (
            ws.cell(row=3, column=1).fill.start_color.rgb
            == EXCEL_FILL_COLOR_NESTED_BLOCK
        )
        assert ws.cell(row=3, column=1).fill.fill_type == "solid"

        # Row 4 (TAG) should have green fill (nested)
        assert (
            ws.cell(row=4, column=1).fill.start_color.rgb
            == EXCEL_FILL_COLOR_NESTED_BLOCK
        )

    def test_format_block_definitions_no_highlighting_non_nested(self) -> None:
        """Test that non-nested blocks are not highlighted."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )
        ws.append(["VALVE", "VALVE", "Inserted", False, "", 8])
        ws.append(["PUMP", "PUMP", "Unused", False, "", 5])

        _format_block_definitions_sheet(wb)

        # Neither row should have green fill
        for row_idx in [2, 3]:
            cell_fill = ws.cell(row=row_idx, column=1).fill
            if cell_fill.fill_type == "solid":
                assert cell_fill.start_color.rgb != EXCEL_FILL_COLOR_NESTED_BLOCK

    def test_format_block_definitions_missing_sheet_handled(self) -> None:
        """Test that missing Block Definitions sheet is handled gracefully."""
        wb = Workbook()
        # Sheet with different name
        ws = cast(Worksheet, wb.active)
        ws.title = "Other Sheet"

        # Should not raise an exception
        _format_block_definitions_sheet(wb)

    def test_format_block_definitions_entire_row_highlighted(self) -> None:
        """Test that entire row (all 6 columns) is highlighted for nested blocks."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )
        ws.append(["PIPE", "PIPE", "Inserted", True, "VALVE", 12])

        _format_block_definitions_sheet(wb)

        # All 6 columns of row 2 should have green fill
        for col_idx in range(1, 7):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_NESTED_BLOCK
            assert cell_fill.fill_type == "solid"

    def test_format_block_definitions_header_text_wrap(self) -> None:
        """Test that header row has text wrap enabled."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )

        _format_block_definitions_sheet(wb)

        # All header cells should have wrap_text enabled
        for cell in ws[1]:
            assert cell.alignment.wrap_text is True

    def test_format_block_definitions_entity_count_right_aligned(self) -> None:
        """Test that entity_count column (F) is right-aligned."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS

        ws.append(
            [
                "block_raw_name",
                "block_resolved_name",
                "block_insertion_status",
                "block_is_nested",
                "block_nested_parent_names",
                "block_entity_count",
            ]
        )
        ws.append(["VALVE", "VALVE", "Inserted", False, "", 8])
        ws.append(["PIPE", "PIPE", "Inserted", True, "VALVE", 12])

        _format_block_definitions_sheet(wb)

        # Column F (entity_count) should be right-aligned for data rows
        assert ws.cell(row=2, column=6).alignment.horizontal == "right"
        assert ws.cell(row=3, column=6).alignment.horizontal == "right"
