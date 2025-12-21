"""
Unit tests for the excel_formatting module - All Blocks sheet.

This test suite validates All Blocks sheet formatting including:
- Auto-filter application
- Column width settings
- Frozen panes
- Three-tier scale highlighting (red, orange, yellow)
- Text wrapping for name and segment columns
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import (
    EXCEL_FILL_COLOR_SCALE_NEGATIVE,
    EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE,
    EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
    EXCEL_SHEET_ALL_BLOCKS,
)
from core.excel_formatting import _format_all_blocks_sheet


class TestAllBlocksFormatting:
    """Test suite for _format_all_blocks_sheet function."""

    def _create_test_workbook_with_headers(self) -> tuple[Workbook, Worksheet]:
        """Create a test workbook with All Blocks sheet and headers."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ALL_BLOCKS

        # Add all 31 column headers (new order from Unit 2)
        headers = [
            "block_raw_name",  # A
            "block_layer_names",  # B (MOVED from position 22)
            "block_resolved_name",  # C
            "block_suggested_trim_left",  # D
            "block_suggested_trim_right",  # E
            "block_suggested_trim_top",  # F
            "block_suggested_trim_bottom",  # G
            "block_vertical_segments",  # H
            "block_horizontal_segments",  # I
            "block_native_width",  # J
            "block_native_height",  # K
            "block_content_zone_detected",  # L
            "block_content_zone_width",  # M
            "block_content_zone_height",  # N
            "block_polygon_count",  # O
            "block_filtered_polygon_count",  # P
            "block_insertion_status",  # Q
            "block_is_nested",  # R
            "block_nested_parent_names",  # S
            "block_entity_count",  # T
            "block_insertion_count",  # U
            "block_layer_count",  # V
            "block_rotation_0",  # W
            "block_rotation_90",  # X
            "block_rotation_180",  # Y
            "block_rotation_270",  # Z
            "block_rotation_other",  # AA
            "block_scale_x",  # AB
            "block_scale_y",  # AC
            "block_attribute_count",  # AD (NEW)
            "block_attribute_data",  # AE (NEW)
        ]
        ws.append(headers)
        return wb, ws

    def _create_test_data_row(
        self,
        raw_name: str = "VALVE",
        layer_names: str = "Layer1, Layer2",
        resolved_name: str = "VALVE",
        scale_x: float | str = 1.0,
        scale_y: float | str = 1.0,
        attribute_count: int = 0,
        attribute_data: str = "",
    ) -> list[object]:
        """Create a test data row with 31 columns in new order."""
        return [
            raw_name,  # A: block_raw_name
            layer_names,  # B: block_layer_names (MOVED)
            resolved_name,  # C: block_resolved_name
            10.0,  # D: trim_left
            10.0,  # E: trim_right
            5.0,  # F: trim_top
            5.0,  # G: trim_bottom
            "10, 80, 10",  # H: vertical_segments
            "5, 40, 5",  # I: horizontal_segments
            100.0,  # J: native_width
            50.0,  # K: native_height
            "TRUE",  # L: content_zone_detected
            80.0,  # M: content_zone_width
            40.0,  # N: content_zone_height
            2,  # O: polygon_count
            2,  # P: filtered_polygon_count
            "Inserted",  # Q: insertion_status
            False,  # R: is_nested
            "",  # S: parent_names
            8,  # T: entity_count
            10,  # U: insertion_count
            2,  # V: layer_count
            5,  # W: rotation_0
            2,  # X: rotation_90
            0,  # Y: rotation_180
            0,  # Z: rotation_270
            0,  # AA: rotation_other
            scale_x,  # AB: scale_x
            scale_y,  # AC: scale_y
            attribute_count,  # AD: attribute_count (NEW)
            attribute_data,  # AE: attribute_data (NEW)
        ]

    def test_format_all_blocks_autofilter(self) -> None:
        """Test that auto-filter is applied to All Blocks sheet."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a data row using the helper method
        ws.append(self._create_test_data_row())

        _format_all_blocks_sheet(wb)

        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:AE2"

    def test_format_all_blocks_column_widths(self) -> None:
        """Test that column widths are set correctly for all 31 columns."""
        wb, ws = self._create_test_workbook_with_headers()

        _format_all_blocks_sheet(wb)

        # Check key column widths (updated for new column order)
        assert ws.column_dimensions["A"].width == 35  # block_raw_name
        assert ws.column_dimensions["B"].width == 35  # block_layer_names (NEW position)
        assert ws.column_dimensions["C"].width == 35  # block_resolved_name
        assert ws.column_dimensions["AB"].width == 12  # block_scale_x
        assert ws.column_dimensions["AC"].width == 12  # block_scale_y
        assert ws.column_dimensions["AD"].width == 12  # block_attribute_count (NEW)
        assert ws.column_dimensions["AE"].width == 60  # block_attribute_data (NEW)
        assert ws.column_dimensions["H"].width == 50  # block_vertical_segments
        assert ws.column_dimensions["I"].width == 50  # block_horizontal_segments
        assert ws.column_dimensions["P"].width == 18  # block_filtered_polygon_count

    def test_format_all_blocks_frozen_panes(self) -> None:
        """Test that frozen panes are applied at B2 (header row and first column)."""
        wb, ws = self._create_test_workbook_with_headers()

        _format_all_blocks_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_format_all_blocks_red_highlighting_varies_negative(self) -> None:
        """Test that rows with VARIES (-) get red highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with VARIES (-) in x_scale (column AB = 28)
        ws.append(self._create_test_data_row(scale_x="VARIES (-)"))

        _format_all_blocks_sheet(wb)

        # Row 2 should have red fill
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
        assert cell_fill.fill_type == "solid"

    def test_format_all_blocks_orange_highlighting_negative_number(self) -> None:
        """Test that rows with single negative scale get orange highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with negative y_scale (column AC = 29)
        ws.append(
            self._create_test_data_row(
                raw_name="PIPE", resolved_name="PIPE", scale_y=-1.0
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should have orange fill
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE
        assert cell_fill.fill_type == "solid"

    def test_format_all_blocks_yellow_highlighting_varies(self) -> None:
        """Test that rows with VARIES (positive only) get yellow highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with VARIES in x_scale (column AB = 28)
        ws.append(
            self._create_test_data_row(
                raw_name="BLOCK1", resolved_name="BLOCK1", scale_x="VARIES"
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should have yellow fill
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
        assert cell_fill.fill_type == "solid"

    def test_format_all_blocks_no_highlighting_positive_scales(self) -> None:
        """Test that rows with positive uniform scales have no highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with positive uniform scales
        ws.append(
            self._create_test_data_row(
                raw_name="BLOCK2",
                resolved_name="BLOCK2",
                layer_names="Layer1",
                scale_x=1.0,
                scale_y=1.5,
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should NOT have any of the scale highlighting colors
        cell_fill = ws.cell(row=2, column=1).fill
        if cell_fill.fill_type == "solid":
            assert cell_fill.start_color.rgb not in [
                EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE,
                EXCEL_FILL_COLOR_SCALE_NEGATIVE,
                EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
            ]

    def test_format_all_blocks_entire_row_highlighted(self) -> None:
        """Test that entire row (all 31 columns) is highlighted for scale issues."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with VARIES (-) to trigger red highlighting
        ws.append(self._create_test_data_row(scale_x="VARIES (-)"))

        _format_all_blocks_sheet(wb)

        # All 31 columns of row 2 should have red fill
        for col_idx in range(1, 32):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
            assert cell_fill.fill_type == "solid"

    def test_format_all_blocks_segment_columns_text_wrap(self) -> None:
        """Test that segment columns (H and I) have text wrapping with top alignment."""
        wb, ws = self._create_test_workbook_with_headers()

        ws.append(self._create_test_data_row())

        _format_all_blocks_sheet(wb)

        # Column H (8 = vertical_segments) should have wrap_text with top vertical alignment
        h_cell = ws.cell(row=2, column=8)
        assert h_cell.alignment.wrap_text is True
        assert h_cell.alignment.vertical == "top"

        # Column I (9 = horizontal_segments) should have wrap_text with top vertical alignment
        i_cell = ws.cell(row=2, column=9)
        assert i_cell.alignment.wrap_text is True
        assert i_cell.alignment.vertical == "top"

    def test_format_all_blocks_header_text_wrap(self) -> None:
        """Test that header row has text wrap enabled."""
        wb, ws = self._create_test_workbook_with_headers()

        _format_all_blocks_sheet(wb)

        # All header cells should have wrap_text enabled
        for cell in ws[1]:
            assert cell.alignment.wrap_text is True

    def test_format_all_blocks_missing_sheet_handled(self) -> None:
        """Test that missing All Blocks sheet is handled gracefully."""
        wb = Workbook()
        # Sheet with different name
        ws = cast(Worksheet, wb.active)
        ws.title = "Other Sheet"

        # Should not raise an exception
        _format_all_blocks_sheet(wb)

    def test_format_all_blocks_highlighting_priority(self) -> None:
        """Test that VARIES (-) has higher priority than VARIES for highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with VARIES (-) in x and VARIES in y
        ws.append(
            self._create_test_data_row(
                raw_name="MIXED",
                resolved_name="MIXED",
                scale_x="VARIES (-)",  # Should trigger red
                scale_y="VARIES",  # Would trigger yellow but red wins
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should have red fill (VARIES (-) takes priority)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

    def test_format_all_blocks_y_scale_varies_negative(self) -> None:
        """Test that VARIES (-) in y_scale also triggers red highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with VARIES (-) in y_scale
        ws.append(
            self._create_test_data_row(
                raw_name="BLOCK_Y", resolved_name="BLOCK_Y", scale_y="VARIES (-)"
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should have red fill
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

    def test_format_all_blocks_negative_x_scale(self) -> None:
        """Test that negative x_scale triggers orange highlighting."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with negative x_scale
        ws.append(
            self._create_test_data_row(
                raw_name="BLOCK_NEG_X", resolved_name="BLOCK_NEG_X", scale_x=-1.0
            )
        )

        _format_all_blocks_sheet(wb)

        # Row 2 should have orange fill
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE

    def test_format_all_blocks_attribute_data_text_wrap(self) -> None:
        """Test that attribute_data column (AE) has text wrapping for multiline content."""
        wb, ws = self._create_test_workbook_with_headers()

        # Add a row with multiline attribute data
        ws.append(
            self._create_test_data_row(
                attribute_count=3,
                attribute_data="DEPT:30\nPROD1:Garage\nPROD2:Doors",
            )
        )

        _format_all_blocks_sheet(wb)

        # Column AE (31 = attribute_data) should have wrap_text with top vertical alignment
        ae_cell = ws.cell(row=2, column=31)
        assert ae_cell.alignment.wrap_text is True
        assert ae_cell.alignment.vertical == "top"

    def test_format_all_blocks_layer_names_text_wrap(self) -> None:
        """Test that layer_names column (B) has text wrapping at new position."""
        wb, ws = self._create_test_workbook_with_headers()

        ws.append(
            self._create_test_data_row(
                layer_names="Layer1, Layer2, Layer3, Layer4",
            )
        )

        _format_all_blocks_sheet(wb)

        # Column B (2 = layer_names) should have wrap_text with top vertical alignment
        b_cell = ws.cell(row=2, column=2)
        assert b_cell.alignment.wrap_text is True
        assert b_cell.alignment.vertical == "top"

    def test_format_all_blocks_attribute_column_widths(self) -> None:
        """Test that attribute columns have correct widths."""
        wb, ws = self._create_test_workbook_with_headers()

        _format_all_blocks_sheet(wb)

        # block_attribute_count (AD) should be 12
        assert ws.column_dimensions["AD"].width == 12
        # block_attribute_data (AE) should be 60
        assert ws.column_dimensions["AE"].width == 60
