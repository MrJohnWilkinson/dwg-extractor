"""
Unit tests for the excel_formatting module - sheet formatting.

This test suite validates sheet-specific formatting functions including:
- Block Analysis sheet formatting
- Layer Analysis sheet formatting
- Entity Summary sheet formatting
- Block Geometry Analysis sheet formatting
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import (
    EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
    EXCEL_SHEET_ANNOTATIONS_ANALYSIS,
    EXCEL_SHEET_BLOCK_ANALYSIS,
    EXCEL_SHEET_BLOCK_DEFINITIONS,
    EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS,
    EXCEL_SHEET_COLOR_ANALYSIS,
    EXCEL_SHEET_ENTITY_SUMMARY,
    EXCEL_SHEET_EXTRACTION_ISSUES,
    EXCEL_SHEET_LAYER_ANALYSIS,
)
from core.excel_formatting import (
    _format_annotations_analysis_sheet,
    _format_block_analysis_sheet,
    _format_block_definitions_sheet,
    _format_block_geometry_analysis_sheet,
    _format_color_analysis_sheet,
    _format_entity_summary_sheet,
    _format_extraction_issues_sheet,
    _format_layer_analysis_sheet,
)


class TestBlockAnalysisFormatting:
    """Test suite for _format_block_analysis_sheet function."""

    def test_format_block_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Block Analysis sheet."""
        # Create test workbook with Block Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS

        # Add headers and sample data
        ws.append(
            [
                "block_name",
                "block_insertion_count",
                "block_entity_count",
                "block_layer_name",
            ]
        )
        ws.append(["VALVE", 10, 8, "Layer1"])

        # Apply formatting
        _format_block_analysis_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:D2"

    def test_format_block_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Block Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS
        ws.append(
            [
                "block_name",
                "block_insertion_count",
                "block_entity_count",
                "block_layer_name",
            ]
        )

        # Apply formatting
        _format_block_analysis_sheet(wb)

        # Verify column widths (4 columns)
        assert ws.column_dimensions["A"].width == 30  # block_name
        assert ws.column_dimensions["B"].width == 25  # block_insertion_count
        assert ws.column_dimensions["C"].width == 25  # block_entity_count
        assert ws.column_dimensions["D"].width == 25  # block_layer_name

    def test_format_block_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Block Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_block_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["B"].width == 25


class TestLayerAnalysisFormatting:
    """Test suite for _format_layer_analysis_sheet function."""

    def test_format_layer_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Layer Analysis sheet."""
        # Create test workbook with Layer Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS

        # Add headers and sample data (5 columns now)
        ws.append(
            [
                "layer_name",
                "layer_block_insertion_count",
                "layer_entity_count",
                "layer_unique_color_count",
                "layer_text_mtext_count",
            ]
        )
        ws.append(["Layer1", 15, 25, 3, 5])

        # Apply formatting
        _format_layer_analysis_sheet(wb)

        # Verify auto-filter is applied (now includes 5 columns)
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:E2"

    def test_format_layer_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Layer Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS
        ws.append(
            [
                "layer_name",
                "layer_block_insertion_count",
                "layer_entity_count",
                "layer_unique_color_count",
                "layer_text_mtext_count",
            ]
        )
        ws.append(["Layer1", 10, 20, 3, 5])
        ws.append(["Layer2", 5, 15, 2, 2])

        # Apply formatting
        _format_layer_analysis_sheet(wb)

        # Verify column widths (5 columns)
        assert ws.column_dimensions["A"].width == 30  # layer_name
        assert ws.column_dimensions["B"].width == 25  # layer_block_insertion_count
        assert ws.column_dimensions["C"].width == 25  # layer_entity_count
        assert ws.column_dimensions["D"].width == 25  # layer_unique_color_count
        assert ws.column_dimensions["E"].width == 25  # layer_text_mtext_count

        # Verify right-alignment on numeric columns (columns B, C, D, E data rows)
        assert ws.cell(row=2, column=2).alignment.horizontal == "right"
        assert ws.cell(row=2, column=3).alignment.horizontal == "right"
        assert ws.cell(row=2, column=4).alignment.horizontal == "right"
        assert ws.cell(row=2, column=5).alignment.horizontal == "right"

    def test_format_layer_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Layer Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_layer_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["B"].width == 25


class TestEntitySummaryFormatting:
    """Test suite for _format_entity_summary_sheet function."""

    def test_format_entity_summary_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Entity Summary sheet."""
        # Create test workbook with Entity Summary sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY

        # Add headers and sample data
        ws.append(["entity_type_name", "entity_type_count"])
        ws.append(["INSERT", 18])

        # Apply formatting
        _format_entity_summary_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:B2"

    def test_format_entity_summary_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Entity Summary sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY
        ws.append(["entity_type_name", "entity_type_count"])

        # Apply formatting
        _format_entity_summary_sheet(wb)

        # Verify column widths (2 columns)
        assert ws.column_dimensions["A"].width == 25  # entity_type_name
        assert ws.column_dimensions["B"].width == 25  # entity_type_count

    def test_format_entity_summary_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Entity Summary sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY

        # Apply formatting (should not crash on empty sheet)
        _format_entity_summary_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 25


class TestBlockGeometryAnalysisFormatting:
    """Test suite for _format_block_geometry_analysis_sheet function."""

    def test_format_geometry_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Block Geometry Analysis sheet."""
        # Create test workbook with Block Geometry Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data (13 columns)
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:M2"

    def test_format_geometry_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that all 13 column widths are set correctly on Block Geometry Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify column widths (13 columns)
        assert ws.column_dimensions["A"].width == 30  # block_name
        assert ws.column_dimensions["B"].width == 25  # block_layer_name
        assert ws.column_dimensions["C"].width == 12  # block_rotation_0
        assert ws.column_dimensions["D"].width == 12  # block_rotation_90
        assert ws.column_dimensions["E"].width == 12  # block_rotation_180
        assert ws.column_dimensions["F"].width == 12  # block_rotation_270
        assert ws.column_dimensions["G"].width == 12  # block_rotation_other
        assert ws.column_dimensions["H"].width == 15  # block_scale_x
        assert ws.column_dimensions["I"].width == 15  # block_scale_y
        assert ws.column_dimensions["J"].width == 20  # block_native_width
        assert ws.column_dimensions["K"].width == 20  # block_native_height
        assert ws.column_dimensions["L"].width == 40  # block_vertical_segments
        assert ws.column_dimensions["M"].width == 40  # block_horizontal_segments

    def test_format_geometry_analysis_yellow_highlighting_varies_x_scale(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in X scale are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in X scale
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                "VARIES",
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            assert cell_fill.fill_type == "solid"

    def test_format_geometry_analysis_yellow_highlighting_varies_y_scale(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in Y scale are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in Y scale
        ws.append(
            [
                "PIPE",
                "Layer1",
                3,
                2,
                0,
                0,
                0,
                1.0,
                "VARIES",
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

    def test_format_geometry_analysis_yellow_highlighting_both_vary(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in both scales are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in both scales
        ws.append(
            ["TAG", "Layer1", 0, 0, 0, 0, 3, "VARIES", "VARIES", 50.0, 25.0, "50", "25"]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

    def test_format_geometry_analysis_no_highlighting_numeric_scales(
        self, temp_dir: str
    ) -> None:
        """Test that rows with numeric scales (no VARIES) are NOT highlighted."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with numeric scales (no variance)
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify NO yellow highlighting on row 2 (data row)
        # Default fill should be PatternFill with no fill_type or None
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            # Check that it's not the yellow fill
            if cell_fill.fill_type == "solid":
                assert (
                    cell_fill.start_color.rgb
                    != EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
                )

    def test_format_geometry_analysis_highlighting_count(self, temp_dir: str) -> None:
        """Test that highlighting count is correct for multiple rows."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add multiple data rows
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )  # No highlight
        ws.append(
            [
                "PIPE",
                "Layer1",
                3,
                2,
                0,
                0,
                0,
                "VARIES",
                1.0,
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )  # Highlight
        ws.append(
            ["TAG", "Layer2", 0, 0, 0, 0, 2, 1.0, "VARIES", 50.0, 25.0, "50", "25"]
        )  # Highlight
        ws.append(
            [
                "DOOR",
                "Layer1",
                10,
                0,
                0,
                0,
                0,
                2.0,
                2.0,
                120.0,
                60.0,
                "15, 90, 15",
                "10, 40, 10",
            ]
        )  # No highlight

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Count highlighted rows (rows 3 and 4 should be highlighted in yellow)
        highlighted_rows = 0
        for row_idx in range(2, 6):  # Rows 2-5 (data rows)
            cell_fill = ws.cell(row=row_idx, column=1).fill
            if (
                cell_fill.fill_type == "solid"
                and cell_fill.start_color.rgb
                == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            ):
                highlighted_rows += 1

        # Should have 2 highlighted rows
        assert highlighted_rows == 2

    def test_format_geometry_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Block Geometry Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_block_geometry_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["H"].width == 15

    def test_format_geometry_analysis_segment_columns_right_aligned(
        self, temp_dir: str
    ) -> None:
        """Test that segment columns L and M are right-aligned in data rows."""
        # Create test workbook with Block Geometry Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data rows (13 columns)
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )
        ws.append(
            [
                "PUMP",
                "Layer2",
                10,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )
        ws.append(
            [
                "TANK",
                "Layer3",
                3,
                1,
                0,
                0,
                0,
                1.0,
                1.0,
                150.0,
                75.0,
                "15, 120, 15",
                "7, 60, 7",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify columns L (12) and M (13) have right-alignment on data rows (rows 2-4)
        for row_idx in range(2, 5):
            l_cell = ws.cell(row=row_idx, column=12)  # Column L
            m_cell = ws.cell(row=row_idx, column=13)  # Column M

            assert l_cell.alignment is not None
            assert l_cell.alignment.horizontal == "right"
            assert m_cell.alignment is not None
            assert m_cell.alignment.horizontal == "right"

        # Verify header row (row 1) maintains wrap_text alignment (not right-aligned)
        header_l = ws.cell(row=1, column=12)
        header_m = ws.cell(row=1, column=13)
        assert header_l.alignment is not None
        assert header_l.alignment.wrap_text is True
        assert header_m.alignment is not None
        assert header_m.alignment.wrap_text is True

        # Verify other columns (e.g., A, B, C) do not have right-alignment
        for row_idx in range(2, 5):
            a_cell = ws.cell(row=row_idx, column=1)  # Column A
            b_cell = ws.cell(row=row_idx, column=2)  # Column B

            # These cells may have alignment set for yellow highlighting, but not right
            if a_cell.alignment is not None:
                assert a_cell.alignment.horizontal != "right"
            if b_cell.alignment is not None:
                assert b_cell.alignment.horizontal != "right"

    def test_format_geometry_analysis_segment_alignment_with_highlighting(
        self, temp_dir: str
    ) -> None:
        """Test that right-alignment works correctly on highlighted rows with scale variance."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data with VARIES in scale columns
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                "VARIES",
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )
        ws.append(
            [
                "PUMP",
                "Layer2",
                10,
                0,
                0,
                0,
                0,
                1.0,
                "VARIES",
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify that rows with VARIES have both yellow fill and right-alignment on columns L and M
        for row_idx in [2, 3]:
            l_cell = ws.cell(row=row_idx, column=12)  # Column L
            m_cell = ws.cell(row=row_idx, column=13)  # Column M

            # Verify yellow fill is applied
            assert l_cell.fill is not None
            assert (
                l_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            )
            assert m_cell.fill is not None
            assert (
                m_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            )

            # Verify right-alignment is also applied
            assert l_cell.alignment is not None
            assert l_cell.alignment.horizontal == "right"
            assert m_cell.alignment is not None
            assert m_cell.alignment.horizontal == "right"


class TestFreezePanes:
    """Test suite for freeze_panes setting across all sheet types."""

    def test_block_analysis_freeze_panes(self, temp_dir: str) -> None:
        """Test that Block Analysis sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS
        ws.append(["block_name", "block_insertion_count"])
        ws.append(["VALVE", 10])

        _format_block_analysis_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_layer_analysis_freeze_panes(self, temp_dir: str) -> None:
        """Test that Layer Analysis sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS
        ws.append(["layer_name", "layer_block_insertion_count"])
        ws.append(["Layer1", 15])

        _format_layer_analysis_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_entity_summary_freeze_panes(self, temp_dir: str) -> None:
        """Test that Entity Summary sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY
        ws.append(["entity_type_name", "entity_type_count"])
        ws.append(["INSERT", 18])

        _format_entity_summary_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_block_geometry_analysis_freeze_panes(self, temp_dir: str) -> None:
        """Test that Block Geometry Analysis sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS
        ws.append(["block_name", "block_layer_name", "block_scale_x"])
        ws.append(["VALVE", "Layer1", 1.0])

        _format_block_geometry_analysis_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_annotations_analysis_freeze_panes(self, temp_dir: str) -> None:
        """Test that Annotations Analysis sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        ws.append(["annotation_contents", "annotation_type"])
        ws.append(["Test text", "TEXT"])

        _format_annotations_analysis_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_color_analysis_freeze_panes(self, temp_dir: str) -> None:
        """Test that Color Analysis sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_COLOR_ANALYSIS
        ws.append(["color_annotation_contents", "color_layer_name"])
        ws.append(["Sample text", "Layer1"])

        _format_color_analysis_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_extraction_issues_freeze_panes(self, temp_dir: str) -> None:
        """Test that Extraction Issues sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_EXTRACTION_ISSUES
        ws.append(["issue_type", "issue_block_name"])
        ws.append(["Missing Handle", "VALVE"])

        _format_extraction_issues_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_block_definitions_freeze_panes(self, temp_dir: str) -> None:
        """Test that Block Definitions sheet freezes first row and first column."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_DEFINITIONS
        ws.append(["block_raw_name", "block_resolved_name"])
        ws.append(["*Model_Space", "Model_Space"])

        _format_block_definitions_sheet(wb)

        assert ws.freeze_panes == "B2"

    def test_all_sheets_freeze_first_row_and_column(self, temp_dir: str) -> None:
        """Test that all eight sheet types freeze both header row and first column (B2)."""
        # Create workbook with all sheet types
        wb = Workbook()

        # Remove default sheet and create all named sheets
        default_sheet = wb.active
        if default_sheet is not None:
            wb.remove(default_sheet)

        # Create and populate all 8 sheets
        sheets_data = [
            (EXCEL_SHEET_BLOCK_ANALYSIS, ["block_name", "count"], ["VALVE", 10]),
            (EXCEL_SHEET_LAYER_ANALYSIS, ["layer_name", "count"], ["Layer1", 15]),
            (EXCEL_SHEET_ENTITY_SUMMARY, ["entity_type_name", "count"], ["INSERT", 18]),
            (
                EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS,
                ["block_name", "scale"],
                ["VALVE", 1.0],
            ),
            (
                EXCEL_SHEET_ANNOTATIONS_ANALYSIS,
                ["annotation_contents", "type"],
                ["Text", "TEXT"],
            ),
            (EXCEL_SHEET_COLOR_ANALYSIS, ["color_contents", "layer"], ["Text", "L1"]),
            (EXCEL_SHEET_EXTRACTION_ISSUES, ["issue_type", "block"], ["Missing", "V"]),
            (EXCEL_SHEET_BLOCK_DEFINITIONS, ["block_raw_name", "resolved"], ["*M", "M"]),
        ]

        for sheet_name, headers, data in sheets_data:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(headers)
            ws.append(data)

        # Apply formatting to all sheets
        _format_block_analysis_sheet(wb)
        _format_layer_analysis_sheet(wb)
        _format_entity_summary_sheet(wb)
        _format_block_geometry_analysis_sheet(wb)
        _format_annotations_analysis_sheet(wb)
        _format_color_analysis_sheet(wb)
        _format_extraction_issues_sheet(wb)
        _format_block_definitions_sheet(wb)

        # Verify freeze_panes = "B2" for all sheets
        expected_freeze = "B2"
        for sheet_name, _, _ in sheets_data:
            ws = wb[sheet_name]
            assert ws.freeze_panes == expected_freeze, (
                f"Sheet '{sheet_name}' has freeze_panes={ws.freeze_panes}, "
                f"expected {expected_freeze}"
            )

    def test_empty_sheets_freeze_panes(self, temp_dir: str) -> None:
        """Test that empty sheets still have freeze_panes set to B2."""
        # Block Analysis
        wb1 = Workbook()
        ws1 = cast(Worksheet, wb1.active)
        ws1.title = EXCEL_SHEET_BLOCK_ANALYSIS
        _format_block_analysis_sheet(wb1)
        assert ws1.freeze_panes == "B2"

        # Entity Summary
        wb2 = Workbook()
        ws2 = cast(Worksheet, wb2.active)
        ws2.title = EXCEL_SHEET_ENTITY_SUMMARY
        _format_entity_summary_sheet(wb2)
        assert ws2.freeze_panes == "B2"

        # Block Definitions
        wb3 = Workbook()
        ws3 = cast(Worksheet, wb3.active)
        ws3.title = EXCEL_SHEET_BLOCK_DEFINITIONS
        _format_block_definitions_sheet(wb3)
        assert ws3.freeze_panes == "B2"
