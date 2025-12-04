"""
Unit tests for the excel_formatting module - negative scale highlighting.

This test suite validates the three-tier highlighting system for negative scales
in the Block Geometry Analysis sheet.
"""

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS
from core.excel_formatting import _format_block_geometry_analysis_sheet


class TestNegativeScaleHighlighting:
    """Test suite for three-tier highlighting system for negative scales."""

    def test_red_highlighting_varies_with_negatives(self) -> None:
        """Test that 'VARIES (-)' gets red fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES (-)"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES (-)")  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
            ), f"Column {col_idx} should have red fill for 'VARIES (-)'"

    def test_orange_highlighting_consistent_negative(self) -> None:
        """Test that -1.0 gets orange fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with negative number
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (negative)
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify orange fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
                f"Column {col_idx} should have orange fill for -1.0"
            )

    def test_yellow_highlighting_varies_positive(self) -> None:
        """Test that 'VARIES' gets yellow fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES" (positive only)
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES")  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            ), f"Column {col_idx} should have yellow fill for 'VARIES'"

    def test_no_highlighting_consistent_positive(self) -> None:
        """Test that 1.0 gets no fill."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with consistent positive values
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=1.0)  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify no highlighting on first column (spot check)
        cell_fill = ws.cell(row=2, column=1).fill
        # Default fill has no color
        assert cell_fill.fill_type is None or cell_fill.start_color.rgb == "00000000", (
            "Row with consistent positive scales should not have highlighting"
        )

    def test_priority_red_over_orange(self) -> None:
        """Test that red (VARIES (-)) has priority over orange (negative number)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both "VARIES (-)" and -1.0
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(
            row=2, column=8, value="VARIES (-)"
        )  # Column H - X scale (red condition)
        ws.cell(row=2, column=9, value=-1.0)  # Column I - Y scale (orange condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill wins (highest priority)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE, (
            "Red should have priority over orange"
        )

    def test_priority_red_over_yellow(self) -> None:
        """Test that red (VARIES (-)) has priority over yellow (VARIES)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both "VARIES (-)" and "VARIES"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(
            row=2, column=8, value="VARIES (-)"
        )  # Column H - X scale (red condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill wins (highest priority)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE, (
            "Red should have priority over yellow"
        )

    def test_priority_orange_over_yellow(self) -> None:
        """Test that orange (negative number) has priority over yellow (VARIES)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both -1.0 and "VARIES"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (orange condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify orange fill wins (higher priority than yellow)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
            "Orange should have priority over yellow"
        )

    def test_entire_row_highlighted(self) -> None:
        """Test that all 18 columns get the same fill when highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers (18 columns)
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
            "Trim L",
            "Trim R",
            "Trim T",
            "Trim B",
            "CZ Detected",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES (-)"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES (-)")
        ws.cell(row=2, column=9, value=1.0)

        _format_block_geometry_analysis_sheet(wb)

        # Verify all 18 columns have red fill
        for col_idx in range(1, 19):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
            ), f"All 18 columns should have red fill, but column {col_idx} doesn't"

    def test_x_scale_triggers_highlight(self) -> None:
        """Test that only X scale having condition triggers row highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row where only X has negative
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (orange condition)
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale (no condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify row is highlighted
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
            "Row should be highlighted when only X scale has condition"
        )

    def test_y_scale_triggers_highlight(self) -> None:
        """Test that only Y scale having condition triggers row highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row where only Y has variance
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=1.0)  # Column H - X scale (no condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify row is highlighted
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE, (
            "Row should be highlighted when only Y scale has condition"
        )
