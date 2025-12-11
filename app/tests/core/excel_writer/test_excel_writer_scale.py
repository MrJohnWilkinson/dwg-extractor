"""
Unit tests for the excel_writer module - negative scale detection and text generation.

This test suite validates:
- Negative scale detection helper functions
- Scale text generation with negative values
- Integration tests using negative_scale_test.dxf asset
"""

import os
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.constants import (
    EXCEL_COLUMN_BLOCK_NAME,
    EXCEL_COLUMN_BLOCK_SCALE_X,
    EXCEL_COLUMN_BLOCK_SCALE_Y,
    EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS,
)
from core.excel_formatting import format_header
from core.excel_writer import write_excel
from core.extractor import ExtractionResult
from core.types import BlockLayerKey


class TestNegativeScaleDetection:
    """Test suite for the _has_negative_scale_in_set helper function."""

    def test_has_negative_scale_x_with_negatives(self) -> None:
        """Test detection of negative X scales."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, 1.0), (-1.0, 1.0)}
        assert _has_negative_scale_in_set(scale_set, "x") is True

    def test_has_negative_scale_x_without_negatives(self) -> None:
        """Test returns False for all positive X scales."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, 1.0), (2.0, 1.0)}
        assert _has_negative_scale_in_set(scale_set, "x") is False

    def test_has_negative_scale_y_with_negatives(self) -> None:
        """Test detection of negative Y scales."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, -1.0), (1.0, -2.0)}
        assert _has_negative_scale_in_set(scale_set, "y") is True

    def test_has_negative_scale_y_without_negatives(self) -> None:
        """Test returns False for all positive Y scales."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, 1.0), (1.0, 2.0)}
        assert _has_negative_scale_in_set(scale_set, "y") is False

    def test_has_negative_scale_mixed_values(self) -> None:
        """Test set with both positive and negative values."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, 1.0), (-1.0, 1.0), (2.0, -2.0)}
        assert _has_negative_scale_in_set(scale_set, "x") is True
        assert _has_negative_scale_in_set(scale_set, "y") is True

    def test_has_negative_scale_empty_set(self) -> None:
        """Test edge case with empty set raises ValueError."""
        from core.excel_writer import _has_negative_scale_in_set

        with pytest.raises(ValueError, match="scale_set cannot be empty"):
            _has_negative_scale_in_set(set(), "x")

    def test_has_negative_scale_invalid_axis(self) -> None:
        """Test error handling for invalid axis parameter."""
        from core.excel_writer import _has_negative_scale_in_set

        scale_set = {(1.0, 1.0)}
        with pytest.raises(ValueError, match="axis must be 'x' or 'y'"):
            _has_negative_scale_in_set(scale_set, "z")


class TestNegativeScaleTextGeneration:
    """Test suite for scale text generation with negative values."""

    def test_scale_text_varies_with_negatives(self, temp_dir: str) -> None:
        """Test that 'VARIES (-)' is displayed when variance exists with negative values."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {
                "TEST": {(1.0, 1.0), (-1.0, 1.0)}
            },  # X varies with negative
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES (-)"

    def test_scale_text_varies_without_negatives(self, temp_dir: str) -> None:
        """Test that 'VARIES' is displayed when variance exists with all positive values."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {
                "TEST": {(1.0, 1.0), (2.0, 1.0)}
            },  # X varies, all positive
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES"

    def test_scale_text_single_negative_value(self, temp_dir: str) -> None:
        """Test that -1.0 is displayed as numeric when scale is consistent negative."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {"TEST": {(-1.0, 1.0)}},  # Consistent negative X scale
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == -1.0

    def test_scale_text_single_positive_value(self, temp_dir: str) -> None:
        """Test that 1.0 is displayed as numeric when scale is consistent positive."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {"TEST": {(1.0, 1.0)}},  # Consistent positive scale
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == 1.0

    def test_scale_text_x_varies_negative_y_consistent(self, temp_dir: str) -> None:
        """Test mixed scenario: X varies with negatives, Y consistent."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {
                "TEST": {(1.0, 2.0), (-1.0, 2.0)}
            },  # X varies with negatives, Y consistent
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES (-)"
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_Y)] == 2.0

    def test_scale_text_both_vary_one_negative(self, temp_dir: str) -> None:
        """Test both axes vary but only one has negatives."""
        data: ExtractionResult = {
            "block_counts": {"TEST": 2},
            "block_entities": {"TEST": 10},
            "block_layer_pairs": {BlockLayerKey(block_name="TEST", layer_name="0"): 2},
            "block_rotation_counts": {},
            "block_scale_data": {
                "TEST": {(1.0, 1.0), (-1.0, 2.0)}
            },  # Both vary, X has negative
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"0": 2},
            "layer_entity_counts": {"0": 20},
            "layer_unique_color_counts": {"0": 0},
            "layer_annotation_counts": {"0": 0},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {
                "TEST": {
                    "native_width": 10.0,
                    "native_height": 10.0,
                    "vertical_segments": [],
                    "horizontal_segments": [],
                }
            },
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
        }

        output_path = os.path.join(temp_dir, "test.dxf")
        excel_path = write_excel(data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES (-)"
        assert df.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_Y)] == "VARIES"


class TestNegativeScaleIntegration:
    """Integration tests using the negative_scale_test.dxf asset."""

    def test_negative_scale_test_dxf_end_to_end(self, temp_dir: str) -> None:
        """Test end-to-end flow with negative_scale_test.dxf verifying all four scenarios."""
        from core.constants import (
            EXCEL_FILL_COLOR_SCALE_NEGATIVE,
            EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE,
            EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
        )
        from core.extractor import extract_blocks

        # Path to test asset
        test_dxf_path = (
            Path(__file__).parent.parent.parent / "assets" / "negative_scale_test.dxf"
        )

        # Extract data from test DXF
        result = extract_blocks(str(test_dxf_path))

        # Verify block_scale_data contains expected scale sets
        assert "VARY_POSITIVE" in result["block_scale_data"]
        assert "MIRROR_CONSISTENT" in result["block_scale_data"]
        assert "VARY_NEGATIVE" in result["block_scale_data"]
        assert "NORMAL" in result["block_scale_data"]

        # Verify scale sets
        vary_positive_scales = result["block_scale_data"]["VARY_POSITIVE"]
        assert (1.0, 1.0) in vary_positive_scales
        assert (2.0, 1.0) in vary_positive_scales

        mirror_scales = result["block_scale_data"]["MIRROR_CONSISTENT"]
        assert (-1.0, 1.0) in mirror_scales

        vary_negative_scales = result["block_scale_data"]["VARY_NEGATIVE"]
        assert (1.0, 1.0) in vary_negative_scales
        assert (-1.0, 1.0) in vary_negative_scales

        normal_scales = result["block_scale_data"]["NORMAL"]
        assert (1.0, 1.0) in normal_scales

        # Generate Excel file
        output_path = os.path.join(temp_dir, "negative_scale_test.dxf")
        excel_path = write_excel(result, output_path)

        # Load Excel and verify scale text values
        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS)

        # Find rows by block name
        vary_positive_row = df[
            df[format_header(EXCEL_COLUMN_BLOCK_NAME)] == "VARY_POSITIVE"
        ].iloc[0]
        mirror_row = df[
            df[format_header(EXCEL_COLUMN_BLOCK_NAME)] == "MIRROR_CONSISTENT"
        ].iloc[0]
        vary_negative_row = df[
            df[format_header(EXCEL_COLUMN_BLOCK_NAME)] == "VARY_NEGATIVE"
        ].iloc[0]
        normal_row = df[df[format_header(EXCEL_COLUMN_BLOCK_NAME)] == "NORMAL"].iloc[0]

        # Verify scale text values
        assert vary_positive_row[format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES"
        assert mirror_row[format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == -1.0
        assert (
            vary_negative_row[format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == "VARIES (-)"
        )
        assert normal_row[format_header(EXCEL_COLUMN_BLOCK_SCALE_X)] == 1.0

        # Load workbook to verify highlighting
        wb = load_workbook(excel_path)
        ws = wb[EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS]

        # Find row numbers for each block (rows are 1-indexed, with header at row 1)
        row_mapping = {}
        for row_idx in range(2, ws.max_row + 1):
            block_name = ws.cell(row=row_idx, column=1).value
            if block_name:
                row_mapping[block_name] = row_idx

        # Verify highlighting colors
        # VARY_POSITIVE should have yellow fill
        vary_positive_row_idx = row_mapping["VARY_POSITIVE"]
        vary_positive_fill = ws.cell(row=vary_positive_row_idx, column=1).fill
        assert (
            vary_positive_fill.start_color.rgb
            == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
        ), "VARY_POSITIVE should have yellow fill"

        # MIRROR_CONSISTENT should have orange fill
        mirror_row_idx = row_mapping["MIRROR_CONSISTENT"]
        mirror_fill = ws.cell(row=mirror_row_idx, column=1).fill
        assert mirror_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
            "MIRROR_CONSISTENT should have orange fill"
        )

        # VARY_NEGATIVE should have red fill
        vary_negative_row_idx = row_mapping["VARY_NEGATIVE"]
        vary_negative_fill = ws.cell(row=vary_negative_row_idx, column=1).fill
        assert (
            vary_negative_fill.start_color.rgb
            == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
        ), "VARY_NEGATIVE should have red fill"

        # NORMAL should have no fill (default)
        normal_row_idx = row_mapping["NORMAL"]
        normal_fill = ws.cell(row=normal_row_idx, column=1).fill
        assert (
            normal_fill.fill_type is None or normal_fill.start_color.rgb == "00000000"
        ), "NORMAL should not have highlighting"
