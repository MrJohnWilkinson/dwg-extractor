"""
Unit tests for the excel_writer module - Block Definitions sheet.

This test suite validates the Block Definitions sheet generation including:
- Sheet creation with correct columns
- Sorting by insertion status and resolved name
- Data population from all_block_definitions
- Empty data handling
- Integration with main write_excel function
"""

import os

import pandas as pd
from openpyxl import load_workbook

from core.constants import (
    EXCEL_COLUMN_BLOCK_ENTITY_COUNT,
    EXCEL_COLUMN_BLOCK_INSERTION_STATUS,
    EXCEL_COLUMN_BLOCK_IS_NESTED,
    EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES,
    EXCEL_COLUMN_BLOCK_RAW_NAME,
    EXCEL_COLUMN_BLOCK_RESOLVED_NAME,
    EXCEL_SHEET_BLOCK_DEFINITIONS,
)
from core.excel_formatting import format_header
from core.excel_writer import write_excel
from core.extractor import ExtractionResult


class TestBlockDefinitionsSheet:
    """Test suite for Block Definitions sheet creation."""

    def test_block_definitions_sheet_exists(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Block Definitions sheet is created when data exists."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        wb = load_workbook(excel_path)
        assert EXCEL_SHEET_BLOCK_DEFINITIONS in wb.sheetnames

    def test_block_definitions_sheet_columns(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Block Definitions sheet has correct 6 columns."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        expected_columns = [
            format_header(EXCEL_COLUMN_BLOCK_RAW_NAME),
            format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME),
            format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS),
            format_header(EXCEL_COLUMN_BLOCK_IS_NESTED),
            format_header(EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES),
            format_header(EXCEL_COLUMN_BLOCK_ENTITY_COUNT),
        ]
        assert list(df.columns) == expected_columns

    def test_block_definitions_sorted_by_status_then_name(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Block Definitions are sorted by status priority, then by resolved name."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # Inserted blocks should come before Nested Only, which comes before System
        statuses = df[format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS)].tolist()

        # Find indices for each status
        inserted_indices = [i for i, s in enumerate(statuses) if s == "Inserted"]
        nested_only_indices = [i for i, s in enumerate(statuses) if s == "Nested Only"]
        system_indices = [i for i, s in enumerate(statuses) if s == "System"]

        # All Inserted should come before Nested Only
        if inserted_indices and nested_only_indices:
            assert max(inserted_indices) < min(nested_only_indices)

        # All Nested Only should come before System
        if nested_only_indices and system_indices:
            assert max(nested_only_indices) < min(system_indices)

    def test_block_definitions_data_populated(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that block definition data is correctly populated."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # Find VALVE row
        valve_rows = df[df[format_header(EXCEL_COLUMN_BLOCK_RAW_NAME)] == "VALVE"]
        assert len(valve_rows) == 1
        valve_row = valve_rows.iloc[0]

        assert valve_row[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"
        assert (
            valve_row[format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS)] == "Inserted"
        )
        assert valve_row[format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == False  # noqa: E712
        assert valve_row[format_header(EXCEL_COLUMN_BLOCK_ENTITY_COUNT)] == 8

    def test_block_definitions_parent_names_formatted(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that parent names are comma-separated strings."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # Find TAG row which has multiple parents
        tag_rows = df[df[format_header(EXCEL_COLUMN_BLOCK_RAW_NAME)] == "TAG"]
        assert len(tag_rows) == 1
        tag_row = tag_rows.iloc[0]

        parent_names = tag_row[format_header(EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES)]
        assert "PIPE" in parent_names
        assert "VALVE" in parent_names
        assert "," in parent_names

    def test_block_definitions_empty_data_no_sheet(self, temp_dir: str) -> None:
        """Test that empty all_block_definitions does not create the sheet."""
        empty_data: ExtractionResult = {
            "block_counts": {},
            "block_entities": {},
            "block_layer_pairs": {},
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {},
            "layer_entity_counts": {},
            "layer_unique_color_counts": {},
            "layer_annotation_counts": {},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
            "block_attribute_data": {},
        }
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(empty_data, output_path)

        wb = load_workbook(excel_path)
        assert EXCEL_SHEET_BLOCK_DEFINITIONS not in wb.sheetnames

    def test_block_definitions_row_count(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that correct number of rows are created."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # sample_extraction_data has 4 block definitions
        assert len(df) == 4

    def test_block_definitions_nested_block_is_nested_true(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that nested blocks have is_nested=True."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # PIPE should have is_nested=True
        pipe_rows = df[df[format_header(EXCEL_COLUMN_BLOCK_RAW_NAME)] == "PIPE"]
        assert len(pipe_rows) == 1
        assert pipe_rows.iloc[0][format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == True  # noqa: E712

        # TAG should have is_nested=True
        tag_rows = df[df[format_header(EXCEL_COLUMN_BLOCK_RAW_NAME)] == "TAG"]
        assert len(tag_rows) == 1
        assert tag_rows.iloc[0][format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == True  # noqa: E712

    def test_block_definitions_system_block_at_end(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that System blocks appear after Inserted and Nested Only blocks."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(sample_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_BLOCK_DEFINITIONS)

        # Get the status column as a list
        statuses = df[format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS)].tolist()

        # System should be the last status in the list
        system_index = None
        for i, s in enumerate(statuses):
            if s == "System":
                system_index = i
                break

        if system_index is not None:
            # Check all non-System statuses appear before System
            for i, s in enumerate(statuses):
                if s != "System":
                    assert i < system_index or statuses[i:].count("System") == 0
