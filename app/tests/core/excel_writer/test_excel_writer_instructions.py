"""
Unit tests for the Instructions sheet creation in excel_writer module.

This test suite validates:
- Instructions sheet is created with correct content
- Instructions sheet appears as the first (leftmost) tab
- Color coding section contains all four colors with descriptions
- Sheet descriptions section contains all 10 sheet descriptions
"""

import os
from pathlib import Path

from openpyxl import load_workbook

from core.constants import EXCEL_SHEET_INSTRUCTIONS
from core.excel_writer import write_excel
from core.extractor import ExtractionResult


class TestInstructionsSheetCreation:
    """Test suite for Instructions sheet creation."""

    def test_instructions_sheet_created(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Instructions sheet is created in the workbook."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        # Create dummy input file
        Path(input_file).touch()

        # Generate Excel
        write_excel(sample_extraction_data, input_file, output_file)

        # Verify Instructions sheet exists
        wb = load_workbook(output_file)
        assert EXCEL_SHEET_INSTRUCTIONS in wb.sheetnames

    def test_instructions_sheet_is_first(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Instructions sheet is the first (leftmost) tab."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        assert wb.sheetnames[0] == EXCEL_SHEET_INSTRUCTIONS

    def test_instructions_sheet_color_coding_section(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Color Coding section header exists."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 2 should be "Color Coding" section header
        assert ws.cell(row=2, column=1).value == "Color Coding"

    def test_instructions_sheet_four_colors(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that all four color descriptions are present."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Check color names in Column A (rows 3-6)
        color_names = [
            ws.cell(row=3, column=1).value,
            ws.cell(row=4, column=1).value,
            ws.cell(row=5, column=1).value,
            ws.cell(row=6, column=1).value,
        ]
        assert "Yellow" in color_names
        assert "Orange" in color_names
        assert "Red" in color_names
        assert "Light Green" in color_names

    def test_instructions_sheet_sheet_descriptions_section(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Sheet Descriptions section header exists."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Row 8 should be "Sheet Descriptions" section header
        assert ws.cell(row=8, column=1).value == "Sheet Descriptions"

    def test_instructions_sheet_ten_sheet_descriptions(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that all 10 sheet descriptions are present."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Sheet names should be in Column A (rows 9-18)
        sheet_names = [ws.cell(row=i, column=1).value for i in range(9, 19)]
        expected_sheets = [
            "All Blocks",
            "Block Analysis",
            "Layer Analysis",
            "Entity Summary",
            "Block Geometry Analysis",
            "Annotations Analysis",
            "Color Analysis",
            "Extraction Issues",
            "Block Definitions",
            "Attribute Analysis",
        ]
        for expected in expected_sheets:
            assert expected in sheet_names, f"Missing sheet description: {expected}"

    def test_instructions_sheet_row_count(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that Instructions sheet has expected number of rows."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        # Expected rows:
        # 1: Header (Identifier, Description)
        # 2: Color Coding section header
        # 3-6: Four color rows
        # 7: Empty separator
        # 8: Sheet Descriptions section header
        # 9-18: Ten sheet description rows
        # Total: 18 rows
        assert ws.max_row == 18

    def test_instructions_sheet_column_headers(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that column headers are set correctly."""
        input_file = os.path.join(temp_dir, "test.dxf")
        output_file = os.path.join(temp_dir, "output.xlsx")

        Path(input_file).touch()
        write_excel(sample_extraction_data, input_file, output_file)

        wb = load_workbook(output_file)
        ws = wb[EXCEL_SHEET_INSTRUCTIONS]

        assert ws.cell(row=1, column=1).value == "Identifier"
        assert ws.cell(row=1, column=2).value == "Description"
