"""
Unit tests for the excel_writer module - All Blocks sheet.

This test suite validates the All Blocks sheet generation including:
- Sheet creation with correct 31 columns
- Row count excluding system blocks
- Sorting by insertion status and resolved name
- Layer count and layer names aggregation
- Rotation counts aggregation across layers
- Scale variance detection
- Geometry data population
- Content zone data population
- Empty data handling
"""

import os

import pandas as pd
from openpyxl import load_workbook

from core.constants import (
    EXCEL_COLUMN_BLOCK_ATTRIBUTE_COUNT,
    EXCEL_COLUMN_BLOCK_ATTRIBUTE_DATA,
    EXCEL_COLUMN_BLOCK_CONTENT_ZONE_DETECTED,
    EXCEL_COLUMN_BLOCK_CONTENT_ZONE_HEIGHT,
    EXCEL_COLUMN_BLOCK_CONTENT_ZONE_WIDTH,
    EXCEL_COLUMN_BLOCK_ENTITY_COUNT,
    EXCEL_COLUMN_BLOCK_FILTERED_POLYGON_COUNT,
    EXCEL_COLUMN_BLOCK_HORIZONTAL_SEGMENTS,
    EXCEL_COLUMN_BLOCK_INSERTION_COUNT,
    EXCEL_COLUMN_BLOCK_INSERTION_STATUS,
    EXCEL_COLUMN_BLOCK_IS_NESTED,
    EXCEL_COLUMN_BLOCK_LAYER_COUNT,
    EXCEL_COLUMN_BLOCK_LAYER_NAMES,
    EXCEL_COLUMN_BLOCK_NATIVE_HEIGHT,
    EXCEL_COLUMN_BLOCK_NATIVE_WIDTH,
    EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES,
    EXCEL_COLUMN_BLOCK_POLYGON_COUNT,
    EXCEL_COLUMN_BLOCK_RAW_NAME,
    EXCEL_COLUMN_BLOCK_RESOLVED_NAME,
    EXCEL_COLUMN_BLOCK_ROTATION_0,
    EXCEL_COLUMN_BLOCK_ROTATION_90,
    EXCEL_COLUMN_BLOCK_ROTATION_180,
    EXCEL_COLUMN_BLOCK_ROTATION_270,
    EXCEL_COLUMN_BLOCK_ROTATION_OTHER,
    EXCEL_COLUMN_BLOCK_SCALE_X,
    EXCEL_COLUMN_BLOCK_SCALE_Y,
    EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_BOTTOM,
    EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_LEFT,
    EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_RIGHT,
    EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_TOP,
    EXCEL_COLUMN_BLOCK_VERTICAL_SEGMENTS,
    EXCEL_SHEET_ALL_BLOCKS,
)
from core.excel_formatting import format_header
from core.excel_writer import _create_all_blocks_sheet
from core.extractor import ExtractionResult
from core.types import BlockLayerKey


class TestAllBlocksSheet:
    """Test suite for All Blocks sheet creation."""

    def test_all_blocks_sheet_created(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that All Blocks sheet is created when non-system blocks exist."""
        import pandas as pd

        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        wb = load_workbook(output_path)
        assert EXCEL_SHEET_ALL_BLOCKS in wb.sheetnames

    def test_all_blocks_sheet_has_31_columns(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that All Blocks sheet has exactly 31 columns."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)
        assert len(df.columns) == 31

    def test_all_blocks_sheet_column_names(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that All Blocks sheet has correct column headers."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        expected_columns = [
            format_header(EXCEL_COLUMN_BLOCK_RAW_NAME),
            format_header(EXCEL_COLUMN_BLOCK_LAYER_NAMES),  # Moved to position B
            format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME),
            format_header(EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_LEFT),
            format_header(EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_RIGHT),
            format_header(EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_TOP),
            format_header(EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_BOTTOM),
            format_header(EXCEL_COLUMN_BLOCK_VERTICAL_SEGMENTS),
            format_header(EXCEL_COLUMN_BLOCK_HORIZONTAL_SEGMENTS),
            format_header(EXCEL_COLUMN_BLOCK_NATIVE_WIDTH),
            format_header(EXCEL_COLUMN_BLOCK_NATIVE_HEIGHT),
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_DETECTED),
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_WIDTH),
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_HEIGHT),
            format_header(EXCEL_COLUMN_BLOCK_POLYGON_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_FILTERED_POLYGON_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS),
            format_header(EXCEL_COLUMN_BLOCK_IS_NESTED),
            format_header(EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES),
            format_header(EXCEL_COLUMN_BLOCK_ENTITY_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_INSERTION_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_LAYER_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_ROTATION_0),
            format_header(EXCEL_COLUMN_BLOCK_ROTATION_90),
            format_header(EXCEL_COLUMN_BLOCK_ROTATION_180),
            format_header(EXCEL_COLUMN_BLOCK_ROTATION_270),
            format_header(EXCEL_COLUMN_BLOCK_ROTATION_OTHER),
            format_header(EXCEL_COLUMN_BLOCK_SCALE_X),
            format_header(EXCEL_COLUMN_BLOCK_SCALE_Y),
            format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_COUNT),
            format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_DATA),
        ]
        assert list(df.columns) == expected_columns

    def test_all_blocks_sheet_excludes_system_blocks(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that System blocks are excluded from All Blocks sheet."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # sample_extraction_data has 4 blocks: VALVE, PIPE, TAG, *Model_Space (System)
        # Only 3 non-system blocks should appear
        assert len(df) == 3

        # Verify no System blocks in the sheet
        statuses = df[format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS)].tolist()
        assert "System" not in statuses

    def test_all_blocks_sheet_sorted_by_status_then_name(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that rows are sorted by insertion status, then by resolved name."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        statuses = df[format_header(EXCEL_COLUMN_BLOCK_INSERTION_STATUS)].tolist()

        # Find indices for each status
        inserted_indices = [i for i, s in enumerate(statuses) if s == "Inserted"]
        nested_only_indices = [i for i, s in enumerate(statuses) if s == "Nested Only"]

        # All Inserted should come before Nested Only
        if inserted_indices and nested_only_indices:
            assert max(inserted_indices) < min(nested_only_indices)

    def test_all_blocks_layer_count_aggregation(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that layer count is correctly aggregated from block_layer_pairs."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE appears on Layer1 and Layer2 (2 layers)
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert len(valve_row) == 1
        assert valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_LAYER_COUNT)] == 2

        # PIPE appears only on Layer1 (1 layer)
        pipe_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "PIPE"]
        assert len(pipe_row) == 1
        assert pipe_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_LAYER_COUNT)] == 1

    def test_all_blocks_layer_names_aggregation(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that layer names are correctly aggregated and sorted."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE appears on Layer1 and Layer2
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        layer_names = valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_LAYER_NAMES)]
        assert "Layer1" in layer_names
        assert "Layer2" in layer_names
        # Should be sorted alphabetically
        assert layer_names == "Layer1, Layer2"

    def test_all_blocks_rotation_counts_aggregated(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that rotation counts are summed across all layers."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE: Layer1 has 5x0deg and 2x90deg, Layer2 has 3x0deg
        # Total: 8x0deg, 2x90deg
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ROTATION_0)] == 8
        assert valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ROTATION_90)] == 2

    def test_all_blocks_insertion_count(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that insertion count is correctly populated from block_counts."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has 10 insertions according to block_counts
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert (
            valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_INSERTION_COUNT)] == 10
        )

    def test_all_blocks_scale_variance_varies(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that scale variance is detected and shows VARIES."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has {(1.0, 1.0), (-1.0, 1.0)} - X varies with negative
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        x_scale = valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_X)]
        assert x_scale == "VARIES (-)"

    def test_all_blocks_scale_single_negative(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that single negative scale shows numeric value."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # PIPE has {(1.0, -1.0)} - Y scale is -1.0
        pipe_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "PIPE"]
        y_scale = pipe_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_SCALE_Y)]
        assert y_scale == -1.0

    def test_all_blocks_geometry_data_populated(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that geometry data is correctly populated."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has native_width=100.0, native_height=50.0
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert (
            valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_NATIVE_WIDTH)] == 100.0
        )
        assert (
            valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_NATIVE_HEIGHT)] == 50.0
        )

    def test_all_blocks_segments_formatted(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that segments are formatted as comma-separated strings."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has vertical_segments=[10.0, 80.0, 10.0]
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        vertical_segments = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_VERTICAL_SEGMENTS)
        ]
        assert vertical_segments == "10, 80, 10"

    def test_all_blocks_content_zone_detected(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that content zone detection is correctly reported."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has content_zone_detected=True
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        detected = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_DETECTED)
        ]
        # Pandas may convert "TRUE"/"FALSE" to bool when reading (use == for numpy bool)
        assert detected == True or detected == "TRUE"  # noqa: E712

        # TAG has content_zone_detected=False
        tag_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "TAG"]
        detected = tag_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_DETECTED)
        ]
        # Pandas may convert "TRUE"/"FALSE" to bool when reading (use == for numpy bool)
        assert detected == False or detected == "FALSE"  # noqa: E712

    def test_all_blocks_content_zone_dimensions(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that content zone dimensions are correctly populated."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has content_zone_width=80.0, content_zone_height=40.0
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        cz_width = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_WIDTH)
        ]
        cz_height = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_HEIGHT)
        ]
        assert cz_width == 80.0
        assert cz_height == 40.0

    def test_all_blocks_trim_values(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that trim values are correctly populated for detected content zones."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has suggested_trim_left=10.0, etc.
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        trim_left = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_SUGGESTED_TRIM_LEFT)
        ]
        assert trim_left == 10.0

    def test_all_blocks_polygon_counts(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that polygon counts are correctly populated."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has polygon_count=2, filtered_polygon_count=2
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        poly_count = valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_POLYGON_COUNT)]
        filtered_count = valve_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_FILTERED_POLYGON_COUNT)
        ]
        assert poly_count == 2
        assert filtered_count == 2

    def test_all_blocks_empty_data_creates_headers_only(self, temp_dir: str) -> None:
        """Test that empty all_block_definitions creates sheet with headers only."""
        empty_data: ExtractionResult = {
            "block_counts": {},
            "block_entities": {},
            "block_layer_pairs": {},
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {},
            "layer_entity_counts": {},
            "layer_unique_color_counts": {},
            "layer_annotation_counts": {},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
            "block_attribute_data": {},
        }
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(empty_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # Should have 31 columns but 0 data rows
        assert len(df.columns) == 31
        assert len(df) == 0

    def test_all_blocks_only_system_blocks_creates_empty_sheet(
        self, temp_dir: str
    ) -> None:
        """Test that when all blocks are system blocks, sheet has headers only."""
        system_only_data: ExtractionResult = {
            "block_counts": {},
            "block_entities": {},
            "block_layer_pairs": {},
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {},
            "layer_entity_counts": {},
            "layer_unique_color_counts": {},
            "layer_annotation_counts": {},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},
            "block_content_zone_data": {},
            "all_block_definitions": {
                "*Model_Space": {
                    "block_raw_name": "*Model_Space",
                    "block_resolved_name": "*Model_Space",
                    "block_insertion_status": "System",
                    "block_is_nested": False,
                    "block_nested_parent_names": [],
                    "block_entity_count": 50,
                },
                "*Paper_Space": {
                    "block_raw_name": "*Paper_Space",
                    "block_resolved_name": "*Paper_Space",
                    "block_insertion_status": "System",
                    "block_is_nested": False,
                    "block_nested_parent_names": [],
                    "block_entity_count": 10,
                },
            },
            "nested_block_parents": {},
            "block_attribute_data": {},
        }
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(system_only_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # Should have 31 columns but 0 data rows (all system blocks excluded)
        assert len(df.columns) == 31
        assert len(df) == 0

    def test_all_blocks_block_with_no_insertions(self, temp_dir: str) -> None:
        """Test that blocks with 0 insertions are handled correctly."""
        data_with_unused_block: ExtractionResult = {
            "block_counts": {},  # Empty - no insertions
            "block_entities": {"UNUSED_BLOCK": 5},
            "block_layer_pairs": {},
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {},
            "layer_entity_counts": {},
            "layer_unique_color_counts": {},
            "layer_annotation_counts": {},
            "annotation_data": {},
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},
            "block_content_zone_data": {},
            "all_block_definitions": {
                "UNUSED_BLOCK": {
                    "block_raw_name": "UNUSED_BLOCK",
                    "block_resolved_name": "UNUSED_BLOCK",
                    "block_insertion_status": "Unused",
                    "block_is_nested": False,
                    "block_nested_parent_names": [],
                    "block_entity_count": 5,
                },
            },
            "nested_block_parents": {},
            "block_attribute_data": {},
        }
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(data_with_unused_block, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # Should have 1 row for UNUSED_BLOCK
        assert len(df) == 1
        row = df.iloc[0]
        assert row[format_header(EXCEL_COLUMN_BLOCK_INSERTION_COUNT)] == 0
        assert row[format_header(EXCEL_COLUMN_BLOCK_LAYER_COUNT)] == 0

    def test_all_blocks_block_without_geometry_data(self, temp_dir: str) -> None:
        """Test that blocks without geometry data have empty geometry fields."""
        data_without_geometry: ExtractionResult = {
            "block_counts": {"BLOCK_NO_GEOMETRY": 5},
            "block_entities": {"BLOCK_NO_GEOMETRY": 3},
            "block_layer_pairs": {
                BlockLayerKey(block_name="BLOCK_NO_GEOMETRY", layer_name="Layer1"): 5
            },
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {"Layer1": 5},
            "layer_entity_counts": {"Layer1": 10},
            "layer_unique_color_counts": {"Layer1": 1},
            "layer_annotation_counts": {"Layer1": 0},
            "annotation_data": {},
            "entity_type_counts": {"INSERT": 5},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},  # No geometry data
            "block_content_zone_data": {},  # No content zone data
            "all_block_definitions": {
                "BLOCK_NO_GEOMETRY": {
                    "block_raw_name": "BLOCK_NO_GEOMETRY",
                    "block_resolved_name": "BLOCK_NO_GEOMETRY",
                    "block_insertion_status": "Inserted",
                    "block_is_nested": False,
                    "block_nested_parent_names": [],
                    "block_entity_count": 3,
                },
            },
            "nested_block_parents": {},
            "block_attribute_data": {},
        }
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(data_without_geometry, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        row = df.iloc[0]
        # Geometry fields should be empty (NaN in pandas)
        assert pd.isna(row[format_header(EXCEL_COLUMN_BLOCK_NATIVE_WIDTH)])
        assert pd.isna(row[format_header(EXCEL_COLUMN_BLOCK_NATIVE_HEIGHT)])
        # Content zone detection should be empty (not detected)
        assert pd.isna(row[format_header(EXCEL_COLUMN_BLOCK_CONTENT_ZONE_DETECTED)])

    def test_all_blocks_nested_block_parent_names(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that nested parent names are formatted as comma-separated strings."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # TAG has nested_parent_names=["VALVE", "PIPE"]
        tag_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "TAG"]
        parent_names = tag_row.iloc[0][
            format_header(EXCEL_COLUMN_BLOCK_NESTED_PARENT_NAMES)
        ]
        assert "VALVE" in parent_names
        assert "PIPE" in parent_names
        assert "," in parent_names

    def test_all_blocks_is_nested_field(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that is_nested field is correctly populated."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE is not nested (use == not 'is' for numpy bool comparison)
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == False  # noqa: E712

        # PIPE and TAG are nested
        pipe_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "PIPE"]
        assert pipe_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == True  # noqa: E712

        tag_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "TAG"]
        assert tag_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_IS_NESTED)] == True  # noqa: E712

    def test_all_blocks_attribute_count(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that attribute count is correctly populated."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE has 3 attributes in fixture
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        assert valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_COUNT)] == 3

        # PIPE has 1 attribute
        pipe_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "PIPE"]
        assert pipe_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_COUNT)] == 1

        # TAG has 0 attributes
        tag_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "TAG"]
        assert tag_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_COUNT)] == 0

    def test_all_blocks_attribute_data_format(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that attribute data is formatted as newline-separated TAG:VALUE pairs."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # VALVE should have formatted attribute data
        valve_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "VALVE"]
        attr_data = valve_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_DATA)]

        # Should contain TAG:VALUE format
        assert "DEPT:30" in attr_data
        assert "PROD1:Garage" in attr_data
        assert "PROD2:Door Openers" in attr_data

        # Should be newline-separated
        assert "\n" in attr_data

    def test_all_blocks_attribute_data_empty_block(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that blocks without attributes have empty attribute data."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # TAG has no attributes
        tag_row = df[df[format_header(EXCEL_COLUMN_BLOCK_RESOLVED_NAME)] == "TAG"]
        attr_data = tag_row.iloc[0][format_header(EXCEL_COLUMN_BLOCK_ATTRIBUTE_DATA)]

        # Should be empty string or NaN
        assert pd.isna(attr_data) or attr_data == ""

    def test_all_blocks_layer_names_column_position(
        self, temp_dir: str, sample_extraction_data: ExtractionResult
    ) -> None:
        """Test that block_layer_names column is at position B (index 1)."""
        output_path = os.path.join(temp_dir, "test_output.xlsx")
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _create_all_blocks_sheet(sample_extraction_data, writer)

        df = pd.read_excel(output_path, sheet_name=EXCEL_SHEET_ALL_BLOCKS)

        # Column at index 1 should be block_layer_names
        assert df.columns[1] == format_header(EXCEL_COLUMN_BLOCK_LAYER_NAMES)
