"""
Unit tests for the excel_writer module - Annotations Analysis sheet.

This test suite validates:
- Annotations Analysis sheet creation
- Column structure and data accuracy
- Sorting and empty data handling
- Layer Analysis integration with annotation counts
"""

import os

import pandas as pd
from openpyxl import load_workbook

from core.constants import (
    EXCEL_COLUMN_LAYER_NAME,
    EXCEL_SHEET_ANNOTATIONS_ANALYSIS,
    EXCEL_SHEET_BLOCK_ANALYSIS,
    EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS,
    EXCEL_SHEET_ENTITY_SUMMARY,
    EXCEL_SHEET_LAYER_ANALYSIS,
)
from core.excel_formatting import format_header
from core.excel_writer import write_excel
from core.extractor import ExtractionResult


class TestAnnotationsAnalysisSheet:
    """Test suite for Annotations Analysis sheet generation."""

    def test_annotations_analysis_sheet_exists(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that Annotations Analysis sheet is created."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        wb = load_workbook(excel_path)
        assert EXCEL_SHEET_ANNOTATIONS_ANALYSIS in wb.sheetnames

    def test_annotations_analysis_sheet_columns(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that Annotations Analysis sheet has correct column structure."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_ANNOTATIONS_ANALYSIS)

        # Verify 8 columns in correct order
        from core.constants import (
            EXCEL_COLUMN_ANNOTATION_COLOR_B,
            EXCEL_COLUMN_ANNOTATION_COLOR_G,
            EXCEL_COLUMN_ANNOTATION_COLOR_R,
            EXCEL_COLUMN_ANNOTATION_COLOR_SAMPLE,
            EXCEL_COLUMN_ANNOTATION_CONTENTS,
            EXCEL_COLUMN_ANNOTATION_COUNT,
            EXCEL_COLUMN_ANNOTATION_LAYER_NAME,
            EXCEL_COLUMN_ANNOTATION_TYPE,
        )

        expected_columns = [
            format_header(EXCEL_COLUMN_ANNOTATION_CONTENTS),
            format_header(EXCEL_COLUMN_ANNOTATION_TYPE),
            format_header(EXCEL_COLUMN_ANNOTATION_LAYER_NAME),
            format_header(EXCEL_COLUMN_ANNOTATION_COLOR_R),
            format_header(EXCEL_COLUMN_ANNOTATION_COLOR_G),
            format_header(EXCEL_COLUMN_ANNOTATION_COLOR_B),
            format_header(EXCEL_COLUMN_ANNOTATION_COLOR_SAMPLE),
            format_header(EXCEL_COLUMN_ANNOTATION_COUNT),
        ]

        assert list(df.columns) == expected_columns
        assert len(df.columns) == 8

    def test_annotations_analysis_sheet_data_accuracy(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that annotation data is correctly written to sheet."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_ANNOTATIONS_ANALYSIS)

        # Verify we have 3 data rows
        assert len(df) == 3

        # Verify data values (check first row which should be sorted by count desc)
        from core.constants import (
            EXCEL_COLUMN_ANNOTATION_COLOR_R,
            EXCEL_COLUMN_ANNOTATION_CONTENTS,
            EXCEL_COLUMN_ANNOTATION_COUNT,
            EXCEL_COLUMN_ANNOTATION_TYPE,
        )

        # First row should be "Sample Text" with count=3 (highest count)
        first_row = df.iloc[0]
        assert (
            first_row[format_header(EXCEL_COLUMN_ANNOTATION_CONTENTS)] == "Sample Text"
        )
        assert first_row[format_header(EXCEL_COLUMN_ANNOTATION_TYPE)] == "TEXT"
        assert first_row[format_header(EXCEL_COLUMN_ANNOTATION_COUNT)] == 3
        assert first_row[format_header(EXCEL_COLUMN_ANNOTATION_COLOR_R)] == 255

    def test_annotations_analysis_sheet_sorting(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that annotations are sorted by count descending."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_ANNOTATIONS_ANALYSIS)

        from core.constants import EXCEL_COLUMN_ANNOTATION_COUNT

        # Verify counts are in descending order
        counts = df[format_header(EXCEL_COLUMN_ANNOTATION_COUNT)].tolist()
        assert counts == sorted(counts, reverse=True)
        assert counts == [3, 2, 1]

    def test_annotations_analysis_sheet_empty_data(self, temp_dir: str) -> None:
        """Test Annotations Analysis sheet with no annotation data."""
        empty_data: ExtractionResult = {
            "block_counts": {},
            "block_entities": {},
            "block_layer_pairs": {},
            "block_rotation_counts": {},
            "block_scale_data": {},
            "block_xdata_apps": {},
            "layer_block_insertion_counts": {},
            "layer_entity_counts": {},
            "layer_unique_color_counts": {},
            "layer_annotation_counts": {},
            "annotation_data": {},  # Empty annotations
            "entity_type_counts": {},
            "color_analysis_data": [],
            "extraction_issues": [],
            "block_trimming_data": {},
            "block_content_zone_data": {},
            "all_block_definitions": {},
            "nested_block_parents": {},
            "block_attribute_data": {},
        }

        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(empty_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_ANNOTATIONS_ANALYSIS)

        # Should have headers but no data rows
        assert len(df) == 0
        assert len(df.columns) == 8

    def test_layer_analysis_uses_annotation_count_column(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that Layer Analysis sheet uses renamed layer_annotation_count column."""
        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_LAYER_ANALYSIS)

        # Verify column header is "Layer Annotation Count" (formatted)
        from core.constants import EXCEL_COLUMN_LAYER_ANNOTATION_COUNT

        expected_header = format_header(EXCEL_COLUMN_LAYER_ANNOTATION_COUNT)
        assert expected_header in df.columns
        assert expected_header == "Layer Annotation Count"

        # Verify data values match layer_annotation_counts from input
        layer1_row = df[df[format_header(EXCEL_COLUMN_LAYER_NAME)] == "Layer1"].iloc[0]
        assert layer1_row[expected_header] == 10

    def test_annotations_analysis_integration_with_real_data(
        self, temp_dir: str
    ) -> None:
        """Integration test with real annotation_test.dxf file."""
        from core.extractor import extract_blocks

        result = extract_blocks("app/tests/assets/annotation_test.dxf")
        output_path = os.path.join(temp_dir, "annotation_test.dxf")
        excel_path = write_excel(result, output_path)

        # Verify sheet exists and has data
        df = pd.read_excel(excel_path, sheet_name=EXCEL_SHEET_ANNOTATIONS_ANALYSIS)
        assert len(df) > 0  # Should have annotation data

        # Verify all columns present
        assert len(df.columns) == 8

        # Verify sorting
        from core.constants import EXCEL_COLUMN_ANNOTATION_COUNT

        counts = df[format_header(EXCEL_COLUMN_ANNOTATION_COUNT)].tolist()
        assert counts == sorted(counts, reverse=True)

    def test_eleven_sheets_created(
        self, temp_dir: str, annotation_extraction_data: ExtractionResult
    ) -> None:
        """Test that exactly 11 sheets are created including Instructions, All Blocks, Color Analysis, Extraction Issues, Block Definitions, and Attribute Analysis."""
        from core.constants import (
            EXCEL_SHEET_ALL_BLOCKS,
            EXCEL_SHEET_ATTRIBUTE_ANALYSIS,
            EXCEL_SHEET_BLOCK_DEFINITIONS,
            EXCEL_SHEET_COLOR_ANALYSIS,
            EXCEL_SHEET_EXTRACTION_ISSUES,
            EXCEL_SHEET_INSTRUCTIONS,
        )

        output_path = os.path.join(temp_dir, "test_drawing.dxf")
        excel_path = write_excel(annotation_extraction_data, output_path)

        wb = load_workbook(excel_path)

        # Verify exactly 11 sheets
        assert len(wb.sheetnames) == 11

        # Verify all expected sheet names
        assert EXCEL_SHEET_INSTRUCTIONS in wb.sheetnames
        assert EXCEL_SHEET_ALL_BLOCKS in wb.sheetnames
        assert EXCEL_SHEET_BLOCK_ANALYSIS in wb.sheetnames
        assert EXCEL_SHEET_LAYER_ANALYSIS in wb.sheetnames
        assert EXCEL_SHEET_ENTITY_SUMMARY in wb.sheetnames
        assert EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS in wb.sheetnames
        assert EXCEL_SHEET_ANNOTATIONS_ANALYSIS in wb.sheetnames
        assert EXCEL_SHEET_COLOR_ANALYSIS in wb.sheetnames
        assert EXCEL_SHEET_EXTRACTION_ISSUES in wb.sheetnames
        assert EXCEL_SHEET_BLOCK_DEFINITIONS in wb.sheetnames
        assert EXCEL_SHEET_ATTRIBUTE_ANALYSIS in wb.sheetnames
