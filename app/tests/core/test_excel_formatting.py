"""
Unit tests for the excel_formatting module.

This test suite validates Excel worksheet formatting utilities including:
- Auto-filter application on all sheets
- Column width settings for each sheet
- Red highlighting for mirrored blocks (negative scales)
- Handling of empty worksheets
- Edge cases and boundary conditions
"""

import tempfile
from typing import Iterator, cast

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.constants import (
    EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE,
    EXCEL_SHEET_BLOCK_ANALYSIS,
    EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS,
    EXCEL_SHEET_ENTITY_SUMMARY,
    EXCEL_SHEET_LAYER_ANALYSIS,
)
from core.excel_formatting import (
    _format_block_analysis_sheet,
    _format_block_geometry_analysis_sheet,
    _format_entity_summary_sheet,
    _format_layer_analysis_sheet,
    format_header,
)


class TestFormatHeader:
    """Test suite for format_header function."""

    def test_format_header_basic(self) -> None:
        """Test simple snake_case conversion."""
        assert format_header("block_name") == "Block Name"

    def test_format_header_multiple_words(self) -> None:
        """Test longer field names."""
        assert (
            format_header("layer_block_insertion_count")
            == "Layer Block Insertion Count"
        )

    def test_format_header_single_word(self) -> None:
        """Test single word."""
        assert format_header("count") == "Count"

    def test_format_header_with_numbers(self) -> None:
        """Test with numbers."""
        assert format_header("rotation_0") == "Rotation 0"
        assert format_header("block_rotation_90") == "Block Rotation 90"
        assert format_header("block_rotation_180") == "Block Rotation 180"

    def test_format_header_empty_string(self) -> None:
        """Test empty string."""
        assert format_header("") == ""

    def test_format_header_no_underscores(self) -> None:
        """Test string without underscores."""
        assert format_header("name") == "Name"

    def test_format_header_acronym_dwg(self) -> None:
        """Test DWG acronym override."""
        assert format_header("dwg_file_name") == "DWG File Name"

    def test_format_header_acronym_dxf(self) -> None:
        """Test DXF acronym override."""
        assert format_header("dxf_version") == "DXF Version"

    def test_format_header_acronym_cad(self) -> None:
        """Test CAD acronym override."""
        assert format_header("cad_entity_type") == "CAD Entity Type"

    def test_format_header_acronym_id(self) -> None:
        """Test ID acronym override."""
        assert format_header("block_id") == "Block ID"

    def test_format_header_entity_type_count(self) -> None:
        """Test entity_type_count conversion."""
        assert format_header("entity_type_count") == "Entity Type Count"

    def test_format_header_block_insertion_count(self) -> None:
        """Test block_insertion_count conversion."""
        assert format_header("block_insertion_count") == "Block Insertion Count"


class TestBlockAnalysisFormatting:
    """Test suite for _format_block_analysis_sheet function."""

    @pytest.fixture
    def temp_dir(self) -> Iterator[str]:
        """Create a temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield tmpdir

    def test_format_block_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Block Analysis sheet."""
        # Create test workbook with Block Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS

        # Add headers and sample data
        ws.append(
            [
                "block_name",
                "block_insertion_count",
                "block_entity_count",
                "block_layer_name",
            ]
        )
        ws.append(["VALVE", 10, 8, "Layer1"])

        # Apply formatting
        _format_block_analysis_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:D2"

    def test_format_block_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Block Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS
        ws.append(
            [
                "block_name",
                "block_insertion_count",
                "block_entity_count",
                "block_layer_name",
            ]
        )

        # Apply formatting
        _format_block_analysis_sheet(wb)

        # Verify column widths (4 columns)
        assert ws.column_dimensions["A"].width == 30  # block_name
        assert ws.column_dimensions["B"].width == 25  # block_insertion_count
        assert ws.column_dimensions["C"].width == 25  # block_entity_count
        assert ws.column_dimensions["D"].width == 25  # block_layer_name

    def test_format_block_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Block Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_block_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["B"].width == 25


class TestLayerAnalysisFormatting:
    """Test suite for _format_layer_analysis_sheet function."""

    @pytest.fixture
    def temp_dir(self) -> Iterator[str]:
        """Create a temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield tmpdir

    def test_format_layer_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Layer Analysis sheet."""
        # Create test workbook with Layer Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS

        # Add headers and sample data (5 columns now)
        ws.append(
            [
                "layer_name",
                "layer_block_insertion_count",
                "layer_entity_count",
                "layer_unique_color_count",
                "layer_text_mtext_count",
            ]
        )
        ws.append(["Layer1", 15, 25, 3, 5])

        # Apply formatting
        _format_layer_analysis_sheet(wb)

        # Verify auto-filter is applied (now includes 5 columns)
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:E2"

    def test_format_layer_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Layer Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS
        ws.append(
            [
                "layer_name",
                "layer_block_insertion_count",
                "layer_entity_count",
                "layer_unique_color_count",
                "layer_text_mtext_count",
            ]
        )
        ws.append(["Layer1", 10, 20, 3, 5])
        ws.append(["Layer2", 5, 15, 2, 2])

        # Apply formatting
        _format_layer_analysis_sheet(wb)

        # Verify column widths (5 columns)
        assert ws.column_dimensions["A"].width == 30  # layer_name
        assert ws.column_dimensions["B"].width == 25  # layer_block_insertion_count
        assert ws.column_dimensions["C"].width == 25  # layer_entity_count
        assert ws.column_dimensions["D"].width == 25  # layer_unique_color_count
        assert ws.column_dimensions["E"].width == 25  # layer_text_mtext_count

        # Verify right-alignment on numeric columns (columns B, C, D, E data rows)
        assert ws.cell(row=2, column=2).alignment.horizontal == "right"
        assert ws.cell(row=2, column=3).alignment.horizontal == "right"
        assert ws.cell(row=2, column=4).alignment.horizontal == "right"
        assert ws.cell(row=2, column=5).alignment.horizontal == "right"

    def test_format_layer_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Layer Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_LAYER_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_layer_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["B"].width == 25


class TestEntitySummaryFormatting:
    """Test suite for _format_entity_summary_sheet function."""

    @pytest.fixture
    def temp_dir(self) -> Iterator[str]:
        """Create a temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield tmpdir

    def test_format_entity_summary_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Entity Summary sheet."""
        # Create test workbook with Entity Summary sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY

        # Add headers and sample data
        ws.append(["entity_type_name", "entity_type_count"])
        ws.append(["INSERT", 18])

        # Apply formatting
        _format_entity_summary_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:B2"

    def test_format_entity_summary_column_widths(self, temp_dir: str) -> None:
        """Test that column widths are set correctly on Entity Summary sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY
        ws.append(["entity_type_name", "entity_type_count"])

        # Apply formatting
        _format_entity_summary_sheet(wb)

        # Verify column widths (2 columns)
        assert ws.column_dimensions["A"].width == 25  # entity_type_name
        assert ws.column_dimensions["B"].width == 25  # entity_type_count

    def test_format_entity_summary_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Entity Summary sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ENTITY_SUMMARY

        # Apply formatting (should not crash on empty sheet)
        _format_entity_summary_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 25


class TestBlockGeometryAnalysisFormatting:
    """Test suite for _format_block_geometry_analysis_sheet function."""

    @pytest.fixture
    def temp_dir(self) -> Iterator[str]:
        """Create a temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield tmpdir

    def test_format_geometry_analysis_autofilter(self, temp_dir: str) -> None:
        """Test that auto-filter is applied to Block Geometry Analysis sheet."""
        # Create test workbook with Block Geometry Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data (13 columns)
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify auto-filter is applied
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref == "A1:M2"

    def test_format_geometry_analysis_column_widths(self, temp_dir: str) -> None:
        """Test that all 13 column widths are set correctly on Block Geometry Analysis sheet."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify column widths (13 columns)
        assert ws.column_dimensions["A"].width == 30  # block_name
        assert ws.column_dimensions["B"].width == 25  # block_layer_name
        assert ws.column_dimensions["C"].width == 12  # block_rotation_0
        assert ws.column_dimensions["D"].width == 12  # block_rotation_90
        assert ws.column_dimensions["E"].width == 12  # block_rotation_180
        assert ws.column_dimensions["F"].width == 12  # block_rotation_270
        assert ws.column_dimensions["G"].width == 12  # block_rotation_other
        assert ws.column_dimensions["H"].width == 15  # block_scale_x
        assert ws.column_dimensions["I"].width == 15  # block_scale_y
        assert ws.column_dimensions["J"].width == 20  # block_native_width
        assert ws.column_dimensions["K"].width == 20  # block_native_height
        assert ws.column_dimensions["L"].width == 40  # block_vertical_segments
        assert ws.column_dimensions["M"].width == 40  # block_horizontal_segments

    def test_format_geometry_analysis_yellow_highlighting_varies_x_scale(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in X scale are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in X scale
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                "VARIES",
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            assert cell_fill.fill_type == "solid"

    def test_format_geometry_analysis_yellow_highlighting_varies_y_scale(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in Y scale are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in Y scale
        ws.append(
            [
                "PIPE",
                "Layer1",
                3,
                2,
                0,
                0,
                0,
                1.0,
                "VARIES",
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

    def test_format_geometry_analysis_yellow_highlighting_both_vary(
        self, temp_dir: str
    ) -> None:
        """Test that rows with VARIES in both scales are highlighted in yellow."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with VARIES in both scales
        ws.append(
            ["TAG", "Layer1", 0, 0, 0, 0, 3, "VARIES", "VARIES", 50.0, 25.0, "50", "25"]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow highlighting on row 2 (data row)
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

    def test_format_geometry_analysis_no_highlighting_numeric_scales(
        self, temp_dir: str
    ) -> None:
        """Test that rows with numeric scales (no VARIES) are NOT highlighted."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add data row with numeric scales (no variance)
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify NO yellow highlighting on row 2 (data row)
        # Default fill should be PatternFill with no fill_type or None
        for col_idx in range(1, 14):  # Columns A through M
            cell_fill = ws.cell(row=2, column=col_idx).fill
            # Check that it's not the yellow fill
            if cell_fill.fill_type == "solid":
                assert (
                    cell_fill.start_color.rgb
                    != EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
                )

    def test_format_geometry_analysis_highlighting_count(self, temp_dir: str) -> None:
        """Test that highlighting count is correct for multiple rows."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )

        # Add multiple data rows
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )  # No highlight
        ws.append(
            [
                "PIPE",
                "Layer1",
                3,
                2,
                0,
                0,
                0,
                "VARIES",
                1.0,
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )  # Highlight
        ws.append(
            ["TAG", "Layer2", 0, 0, 0, 0, 2, 1.0, "VARIES", 50.0, 25.0, "50", "25"]
        )  # Highlight
        ws.append(
            [
                "DOOR",
                "Layer1",
                10,
                0,
                0,
                0,
                0,
                2.0,
                2.0,
                120.0,
                60.0,
                "15, 90, 15",
                "10, 40, 10",
            ]
        )  # No highlight

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Count highlighted rows (rows 3 and 4 should be highlighted in yellow)
        highlighted_rows = 0
        for row_idx in range(2, 6):  # Rows 2-5 (data rows)
            cell_fill = ws.cell(row=row_idx, column=1).fill
            if (
                cell_fill.fill_type == "solid"
                and cell_fill.start_color.rgb
                == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            ):
                highlighted_rows += 1

        # Should have 2 highlighted rows
        assert highlighted_rows == 2

    def test_format_geometry_analysis_empty_sheet(self, temp_dir: str) -> None:
        """Test that empty Block Geometry Analysis sheet is handled gracefully."""
        # Create empty workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Apply formatting (should not crash on empty sheet)
        _format_block_geometry_analysis_sheet(wb)

        # Column widths should still be set
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["H"].width == 15

    def test_format_geometry_analysis_segment_columns_right_aligned(
        self, temp_dir: str
    ) -> None:
        """Test that segment columns L and M are right-aligned in data rows."""
        # Create test workbook with Block Geometry Analysis sheet
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data rows (13 columns)
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                1.0,
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )
        ws.append(
            [
                "PUMP",
                "Layer2",
                10,
                0,
                0,
                0,
                0,
                1.0,
                1.0,
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )
        ws.append(
            [
                "TANK",
                "Layer3",
                3,
                1,
                0,
                0,
                0,
                1.0,
                1.0,
                150.0,
                75.0,
                "15, 120, 15",
                "7, 60, 7",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify columns L (12) and M (13) have right-alignment on data rows (rows 2-4)
        for row_idx in range(2, 5):
            l_cell = ws.cell(row=row_idx, column=12)  # Column L
            m_cell = ws.cell(row=row_idx, column=13)  # Column M

            assert l_cell.alignment is not None
            assert l_cell.alignment.horizontal == "right"
            assert m_cell.alignment is not None
            assert m_cell.alignment.horizontal == "right"

        # Verify header row (row 1) maintains wrap_text alignment (not right-aligned)
        header_l = ws.cell(row=1, column=12)
        header_m = ws.cell(row=1, column=13)
        assert header_l.alignment is not None
        assert header_l.alignment.wrap_text is True
        assert header_m.alignment is not None
        assert header_m.alignment.wrap_text is True

        # Verify other columns (e.g., A, B, C) do not have right-alignment
        for row_idx in range(2, 5):
            a_cell = ws.cell(row=row_idx, column=1)  # Column A
            b_cell = ws.cell(row=row_idx, column=2)  # Column B

            # These cells may have alignment set for yellow highlighting, but not right
            if a_cell.alignment is not None:
                assert a_cell.alignment.horizontal != "right"
            if b_cell.alignment is not None:
                assert b_cell.alignment.horizontal != "right"

    def test_format_geometry_analysis_segment_alignment_with_highlighting(
        self, temp_dir: str
    ) -> None:
        """Test that right-alignment works correctly on highlighted rows with scale variance."""
        # Create test workbook
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers and sample data with VARIES in scale columns
        ws.append(
            [
                "block_name",
                "block_layer_name",
                "block_rotation_0",
                "block_rotation_90",
                "block_rotation_180",
                "block_rotation_270",
                "block_rotation_other",
                "block_scale_x",
                "block_scale_y",
                "block_native_width",
                "block_native_height",
                "block_vertical_segments",
                "block_horizontal_segments",
            ]
        )
        ws.append(
            [
                "VALVE",
                "Layer1",
                5,
                2,
                0,
                0,
                0,
                "VARIES",
                1.0,
                100.0,
                50.0,
                "10, 80, 10",
                "5, 40, 5",
            ]
        )
        ws.append(
            [
                "PUMP",
                "Layer2",
                10,
                0,
                0,
                0,
                0,
                1.0,
                "VARIES",
                200.0,
                100.0,
                "20, 160, 20",
                "10, 80, 10",
            ]
        )

        # Apply formatting
        _format_block_geometry_analysis_sheet(wb)

        # Verify that rows with VARIES have both yellow fill and right-alignment on columns L and M
        for row_idx in [2, 3]:
            l_cell = ws.cell(row=row_idx, column=12)  # Column L
            m_cell = ws.cell(row=row_idx, column=13)  # Column M

            # Verify yellow fill is applied
            assert l_cell.fill is not None
            assert (
                l_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            )
            assert m_cell.fill is not None
            assert (
                m_cell.fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            )

            # Verify right-alignment is also applied
            assert l_cell.alignment is not None
            assert l_cell.alignment.horizontal == "right"
            assert m_cell.alignment is not None
            assert m_cell.alignment.horizontal == "right"


class TestNegativeScaleHighlighting:
    """Test suite for three-tier highlighting system for negative scales."""

    @pytest.fixture
    def temp_dir(self) -> Iterator[str]:
        """Create a temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield tmpdir

    def test_red_highlighting_varies_with_negatives(self) -> None:
        """Test that 'VARIES (-)' gets red fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES (-)"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES (-)")  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
            ), f"Column {col_idx} should have red fill for 'VARIES (-)'"

    def test_orange_highlighting_consistent_negative(self) -> None:
        """Test that -1.0 gets orange fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with negative number
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (negative)
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify orange fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
                f"Column {col_idx} should have orange fill for -1.0"
            )

    def test_yellow_highlighting_varies_positive(self) -> None:
        """Test that 'VARIES' gets yellow fill."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES" (positive only)
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES")  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify yellow fill on all columns
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE
            ), f"Column {col_idx} should have yellow fill for 'VARIES'"

    def test_no_highlighting_consistent_positive(self) -> None:
        """Test that 1.0 gets no fill."""
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with consistent positive values
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=1.0)  # Column H - X scale
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale

        _format_block_geometry_analysis_sheet(wb)

        # Verify no highlighting on first column (spot check)
        cell_fill = ws.cell(row=2, column=1).fill
        # Default fill has no color
        assert cell_fill.fill_type is None or cell_fill.start_color.rgb == "00000000", (
            "Row with consistent positive scales should not have highlighting"
        )

    def test_priority_red_over_orange(self) -> None:
        """Test that red (VARIES (-)) has priority over orange (negative number)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both "VARIES (-)" and -1.0
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(
            row=2, column=8, value="VARIES (-)"
        )  # Column H - X scale (red condition)
        ws.cell(row=2, column=9, value=-1.0)  # Column I - Y scale (orange condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill wins (highest priority)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE, (
            "Red should have priority over orange"
        )

    def test_priority_red_over_yellow(self) -> None:
        """Test that red (VARIES (-)) has priority over yellow (VARIES)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both "VARIES (-)" and "VARIES"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(
            row=2, column=8, value="VARIES (-)"
        )  # Column H - X scale (red condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify red fill wins (highest priority)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE, (
            "Red should have priority over yellow"
        )

    def test_priority_orange_over_yellow(self) -> None:
        """Test that orange (negative number) has priority over yellow (VARIES)."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with both -1.0 and "VARIES"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (orange condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify orange fill wins (higher priority than yellow)
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
            "Orange should have priority over yellow"
        )

    def test_entire_row_highlighted(self) -> None:
        """Test that all 13 columns get the same fill when highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with "VARIES (-)"
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value="VARIES (-)")
        ws.cell(row=2, column=9, value=1.0)

        _format_block_geometry_analysis_sheet(wb)

        # Verify all 13 columns have red fill
        for col_idx in range(1, 14):
            cell_fill = ws.cell(row=2, column=col_idx).fill
            assert (
                cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_NEGATIVE
            ), f"All 13 columns should have red fill, but column {col_idx} doesn't"

    def test_x_scale_triggers_highlight(self) -> None:
        """Test that only X scale having condition triggers row highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_NEGATIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row where only X has negative
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=-1.0)  # Column H - X scale (orange condition)
        ws.cell(row=2, column=9, value=1.0)  # Column I - Y scale (no condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify row is highlighted
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_NEGATIVE, (
            "Row should be highlighted when only X scale has condition"
        )

    def test_y_scale_triggers_highlight(self) -> None:
        """Test that only Y scale having condition triggers row highlighting."""
        from core.constants import EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_BLOCK_GEOMETRY_ANALYSIS

        # Add headers
        headers = [
            "Block Name",
            "Layer",
            "R0",
            "R90",
            "R180",
            "R270",
            "R Other",
            "Scale X",
            "Scale Y",
            "Width",
            "Height",
            "V Segs",
            "H Segs",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row where only Y has variance
        ws.cell(row=2, column=1, value="TEST_BLOCK")
        ws.cell(row=2, column=8, value=1.0)  # Column H - X scale (no condition)
        ws.cell(
            row=2, column=9, value="VARIES"
        )  # Column I - Y scale (yellow condition)

        _format_block_geometry_analysis_sheet(wb)

        # Verify row is highlighted
        cell_fill = ws.cell(row=2, column=1).fill
        assert cell_fill.start_color.rgb == EXCEL_FILL_COLOR_SCALE_VARIANCE_POSITIVE, (
            "Row should be highlighted when only Y scale has condition"
        )


class TestAnnotationsAnalysisFormatting:
    """Test suite for _format_annotations_analysis_sheet function."""

    def test_annotations_analysis_auto_filter(self) -> None:
        """Test that auto-filter is applied to Annotations Analysis sheet."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers and data
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row
        ws.cell(row=2, column=1, value="Sample Text")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=255)
        ws.cell(row=2, column=5, value=0)
        ws.cell(row=2, column=6, value=0)
        ws.cell(row=2, column=7, value="")
        ws.cell(row=2, column=8, value=3)

        _format_annotations_analysis_sheet(wb)

        # Verify auto-filter
        assert ws.auto_filter is not None
        assert ws.auto_filter.ref == "A1:H2"

    def test_annotations_analysis_frozen_panes(self) -> None:
        """Test that frozen panes are applied at A2."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        _format_annotations_analysis_sheet(wb)

        # Verify frozen panes
        assert ws.freeze_panes is not None
        assert ws.freeze_panes == "A2"

    def test_annotations_analysis_column_widths(self) -> None:
        """Test that column widths are correctly set."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        _format_annotations_analysis_sheet(wb)

        # Verify column widths
        assert ws.column_dimensions["A"].width == 60  # annotation_contents
        assert ws.column_dimensions["B"].width == 15  # annotation_type
        assert ws.column_dimensions["C"].width == 25  # annotation_layer_name
        assert ws.column_dimensions["D"].width == 12  # annotation_color_r
        assert ws.column_dimensions["E"].width == 12  # annotation_color_g
        assert ws.column_dimensions["F"].width == 12  # annotation_color_b
        assert ws.column_dimensions["G"].width == 12  # annotation_color_sample
        assert ws.column_dimensions["H"].width == 20  # annotation_count

    def test_annotations_analysis_text_wrapping(self) -> None:
        """Test that text wrapping is enabled on annotation_contents column."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers and data
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data rows
        ws.cell(row=2, column=1, value="Long text content that should wrap")
        ws.cell(row=3, column=1, value="Another long text")

        _format_annotations_analysis_sheet(wb)

        # Verify text wrapping on column A (annotation_contents)
        assert ws.cell(row=2, column=1).alignment.wrap_text is True
        assert ws.cell(row=3, column=1).alignment.wrap_text is True

    def test_annotations_analysis_color_fill(self) -> None:
        """Test that RGB color fills are applied to annotation_color_sample column."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with RGB values
        ws.cell(row=2, column=1, value="Red Text")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=255)  # R
        ws.cell(row=2, column=5, value=0)  # G
        ws.cell(row=2, column=6, value=0)  # B
        ws.cell(row=2, column=7, value="")  # color_sample
        ws.cell(row=2, column=8, value=1)

        # Add another row with different color
        ws.cell(row=3, column=1, value="Green Text")
        ws.cell(row=3, column=2, value="MTEXT")
        ws.cell(row=3, column=3, value="Layer2")
        ws.cell(row=3, column=4, value=0)  # R
        ws.cell(row=3, column=5, value=255)  # G
        ws.cell(row=3, column=6, value=0)  # B
        ws.cell(row=3, column=7, value="")  # color_sample
        ws.cell(row=3, column=8, value=2)

        _format_annotations_analysis_sheet(wb)

        # Verify color fills on column G (annotation_color_sample)
        # Row 2 should have red fill (FF0000)
        cell_fill_2 = ws.cell(row=2, column=7).fill
        assert cell_fill_2.start_color.rgb == "00FF0000"  # openpyxl uses ARGB format

        # Row 3 should have green fill (00FF00)
        cell_fill_3 = ws.cell(row=3, column=7).fill
        assert cell_fill_3.start_color.rgb == "0000FF00"

    def test_annotations_analysis_empty_sheet(self) -> None:
        """Test formatting with headers only (no data rows)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add only headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Should not raise exception
        _format_annotations_analysis_sheet(wb)

        # Verify basic formatting still applied
        assert ws.freeze_panes == "A2"
        assert ws.column_dimensions["A"].width == 60

    def test_annotations_analysis_invalid_rgb_values(self) -> None:
        """Test handling of invalid RGB values (out of range)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add data row with invalid RGB values (should be clamped or skipped)
        ws.cell(row=2, column=1, value="Invalid Color")
        ws.cell(row=2, column=2, value="TEXT")
        ws.cell(row=2, column=3, value="Layer1")
        ws.cell(row=2, column=4, value=300)  # Invalid R (> 255)
        ws.cell(row=2, column=5, value=-10)  # Invalid G (< 0)
        ws.cell(row=2, column=6, value=100)  # Valid B
        ws.cell(row=2, column=7, value="")
        ws.cell(row=2, column=8, value=1)

        # Should not raise exception
        _format_annotations_analysis_sheet(wb)

    def test_annotations_analysis_rgb_boundary_values(self) -> None:
        """Test RGB color fills with boundary values (0 and 255)."""
        from core.constants import EXCEL_SHEET_ANNOTATIONS_ANALYSIS
        from core.excel_formatting import _format_annotations_analysis_sheet

        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = EXCEL_SHEET_ANNOTATIONS_ANALYSIS

        # Add headers
        headers = [
            "Annotation Contents",
            "Annotation Type",
            "Annotation Layer Name",
            "Annotation Color R",
            "Annotation Color G",
            "Annotation Color B",
            "Annotation Color Sample",
            "Annotation Count",
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Add row with white color (255, 255, 255)
        ws.cell(row=2, column=1, value="White Text")
        ws.cell(row=2, column=4, value=255)
        ws.cell(row=2, column=5, value=255)
        ws.cell(row=2, column=6, value=255)
        ws.cell(row=2, column=7, value="")

        # Add row with black color (0, 0, 0)
        ws.cell(row=3, column=1, value="Black Text")
        ws.cell(row=3, column=4, value=0)
        ws.cell(row=3, column=5, value=0)
        ws.cell(row=3, column=6, value=0)
        ws.cell(row=3, column=7, value="")

        _format_annotations_analysis_sheet(wb)

        # Verify white fill
        white_fill = ws.cell(row=2, column=7).fill
        assert white_fill.start_color.rgb == "00FFFFFF"

        # Verify black fill
        black_fill = ws.cell(row=3, column=7).fill
        assert black_fill.start_color.rgb == "00000000"
